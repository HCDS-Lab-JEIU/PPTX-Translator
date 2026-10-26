"""Document-wide image screening and safe optional text compression."""
from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from PIL import Image
from pptx import Presentation
from pptx.util import Cm
from pptx_styles import styled_frame

from main import extract_pptx
from pptx_pipeline.ai import AIResponse
from pptx_pipeline.layout import (LayoutPolicy, PPTXLayoutModifier, inspect_layout,
                                  apply_layout_changes, text_and_math_signature)
from pptx_pipeline.layout_overview import damage_pages, metric_candidates
from pptx_pipeline.layout_text import validate_compression
from pptx_gui.engine import Stopped
from tests.test_pipeline import FakeRenderer
from tests.test_native_math import native_equation


class OverviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.pptx'
        prs = Presentation()
        for _ in range(12):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Cm(2), Cm(3), Cm(15), Cm(3)).text = 'The current value is 10 mA'
        prs.save(self.source)
        self.manifest = self.root / 'content.json'
        self.document = extract_pptx(self.source, self.root / 'extract')
        self.document['translate_all'] = True
        self.manifest.write_text(json.dumps(self.document), encoding='utf-8')

    def provider(self, selected=(), compression=None, reject_screen=False):
        class Provider:
            model = 'fake'
            def __init__(inner): inner.calls = []; inner.screen_calls = 0
            def generate(inner, prompt, *, images=()):
                inner.calls.append((prompt, list(images)))
                if 'Screen the whole presentation' in prompt:
                    inner.screen_calls += 1
                    if reject_screen:
                        return AIResponse('{"reviews":[]}', tuple(images))
                    pages = json.loads(prompt.split('Do not use tools or modify files.\n')[1])['pages']
                    return AIResponse(json.dumps({'reviews':[{'page': p, 'needs_changes': p in selected,
                        'issues': ['overlap'] if p in selected else [],
                        'defects': [{'kind':'overlap','evidence':'New body text obscures a diagram label'}]
                            if p in selected else []} for p in pages]}), tuple(images))
                data = json.loads(prompt[prompt.index('{"images":'):])
                pages = [slide['page'] for slide in data['slides']]
                editing = not any('after_' in str(image) for image in images)
                changes = []
                if editing:
                    for p in pages:
                        changes.append({'page': p, 'shape_id': 2, 'reason': 'fit without changing meaning',
                                        **({'text': compression} if compression else {'font_scale': 0.8})})
                return AIResponse(json.dumps({'reviews':[{'page': p, 'needs_changes': editing,
                    'issues': ['wrap'] if editing else []} for p in pages], 'changes': changes}), tuple(images))
        return Provider()

    def modify(self, provider, **kwargs):
        renderer = FakeRenderer()
        report = PPTXLayoutModifier(provider, renderer, policy=kwargs.pop('policy', LayoutPolicy(max_passes=1)),
                                    **kwargs).modify(self.source, self.root / 'result.pptx',
            reference=self.source, manifest=self.manifest, workdir=self.root / 'work')
        return report, renderer

    def test_clean_document_has_one_overview_request_and_no_repair(self):
        provider = self.provider()
        report, renderer = self.modify(provider)
        self.assertEqual(len(provider.calls), 1)
        self.assertEqual(len(provider.calls[0][1]), 3)  # four page pairs per JPEG
        self.assertEqual(renderer.requested_pages, [list(range(1, 13))] * 2)
        self.assertEqual(report['screening']['selected_pages'], [])
        self.assertEqual(report['status'], 'verified')
        self.assertEqual((self.root / 'result.pptx').read_bytes(), self.source.read_bytes())
        with Image.open(provider.calls[0][1][0]) as sheet:
            self.assertEqual(sheet.size, (1280, 1560))

    def test_only_selected_pages_are_detailed_and_rerendered(self):
        provider = self.provider(selected=(2, 9))
        report, renderer = self.modify(provider)
        self.assertEqual(len(provider.calls), 3)  # overview, repair, verification
        self.assertEqual([len(images) for _, images in provider.calls], [3, 4, 4])
        self.assertEqual(renderer.requested_pages[-1], [2, 9])
        self.assertEqual(report['screening']['selected_pages'], [2, 9])
        self.assertEqual(report['status'], 'verified')
        self.assertEqual(report['reviewed_pages'], list(range(1, 13)))

    def test_failed_screen_marks_unverified_without_repairing_entire_deck(self):
        provider = self.provider(reject_screen=True)
        report, renderer = self.modify(provider)
        self.assertTrue(report['screening']['failed'])
        self.assertEqual(len(renderer.calls), 2)
        self.assertEqual(report['reviewed_pages'], [])
        self.assertEqual(report['unreviewed_pages'], list(range(1, 13)))
        self.assertEqual(report['screening']['selected_pages'], [])
        self.assertEqual(report['status'], 'needs_review')

    def test_one_line_to_two_lines_is_only_advisory_without_overflow(self):
        renderer = FakeRenderer()
        before = renderer.render(self.source, self.root / 'before')
        after = deepcopy(before)
        before.slides[4]['text_metrics'][0].update(line_count=1, overflow=False)
        after.slides[4]['text_metrics'][0].update(line_count=2, overflow=False)
        candidates = metric_candidates(before, after, list(range(1, 13)))
        self.assertEqual(list(candidates), [4])
        self.assertEqual(candidates[4][0]['severity'], 'advisory')

    def test_optional_compression_is_applied_and_reported(self):
        provider = self.provider(selected=(2,), compression='Current is 10 mA')
        report, renderer = self.modify(provider, policy=LayoutPolicy(max_passes=1, allow_text_compression=True))
        prs = Presentation(self.root / 'result.pptx')
        self.assertEqual(prs.slides[1].shapes[0].text, 'Current is 10 mA')
        self.assertEqual(prs.slides[0].shapes[0].text, 'The current value is 10 mA')
        self.assertEqual(report['status'], 'needs_review')
        self.assertEqual([c['page'] for c in report['text_compressions']], [2])
        self.assertEqual(renderer.requested_pages[-1], [2])

    def test_compression_rejects_number_unit_negation_style_and_code_changes(self):
        frame = {'text': 'The current value is not 10 mA', 'text_editable': True}
        for text in ('Current is not 20 mA', 'Current is not 10 A', 'Current is 10 mA'):
            with self.subTest(text=text), self.assertRaises(ValueError):
                validate_compression(frame, text)
        with self.assertRaises(ValueError):
            validate_compression({'text': 'return current_value + 10', 'text_editable': True}, 'current + 10')
        with self.assertRaises(ValueError):
            validate_compression({'text':'Long voltage text', 'text_editable': True,
                'styled_text':'[[STYLE_0]]Long [[/STYLE_0]][[STYLE_1]]voltage text[[/STYLE_1]]'}, 'Voltage')

    def test_compression_preserves_native_equation_xml(self):
        prs = Presentation(self.source)
        frame = prs.slides[0].shapes[0].text_frame
        frame.paragraphs[0]._p.append(native_equation())
        prs.save(self.source)
        document = extract_pptx(self.source, self.root / 'math-extract')
        document['translate_all'] = True
        renderer = FakeRenderer()
        render = renderer.render(self.source, self.root / 'math-render')
        policy = LayoutPolicy(allow_text_compression=True)
        snapshots = inspect_layout(self.source, render, document, policy)
        info = snapshots[1]['shapes'][0]['frames'][0]
        shortened = info['text'].replace('The current value is', 'Current is')
        plan = {'reviews':[{'page':1,'needs_changes':True,'issues':['wrap']}],
                'changes':[{'page':1,'shape_id':2,'text':shortened,'reason':'fit'}]}
        apply_layout_changes(self.source, self.root / 'math-result.pptx', plan, {1:snapshots[1]}, policy, snapshots)
        before = text_and_math_signature(Presentation(self.source))[1]
        after = text_and_math_signature(Presentation(self.root / 'math-result.pptx'))[1]
        self.assertEqual(before, after)

    def test_resume_reuses_render_screen_and_completed_detailed_batches(self):
        provider = self.provider(selected=tuple(range(1, 13)))
        original_generate = provider.generate
        attempted = []
        def interrupted(*args, **kwargs):
            attempted.append(1)
            if len(attempted) == 3:
                raise Stopped()
            return original_generate(*args, **kwargs)
        provider.generate = interrupted
        renderer = FakeRenderer()
        modifier = PPTXLayoutModifier(provider, renderer, policy=LayoutPolicy(max_passes=1))
        options = dict(reference=self.source, manifest=self.manifest, workdir=self.root / 'resume')
        with self.assertRaises(Stopped):
            modifier.modify(self.source, self.root / 'resume-result.pptx', **options)
        report = modifier.modify(self.source, self.root / 'resume-result.pptx', **options)
        self.assertEqual(provider.screen_calls, 1)
        self.assertEqual(len(renderer.calls), 3)
        first_batch = [prompt for prompt, images in provider.calls if 'Review actual slide' in prompt
                       and 'before' in str(images[-1]) and 'slide_006' in str(images[-1])]
        self.assertEqual(len(first_batch), 1)
        self.assertEqual(report['status'], 'verified')

    def test_mixed_style_table_compression_preserves_emphasis(self):
        prs = Presentation(self.source)
        table = prs.slides[0].shapes.add_table(1, 1, Cm(2), Cm(8), Cm(15), Cm(3))
        frame = table.table.cell(0, 0).text_frame
        paragraph = frame.paragraphs[0]
        paragraph.add_run().text = 'The current value is '
        emphasis = paragraph.add_run()
        emphasis.text = '10 mA'
        emphasis.font.bold = True
        prs.save(self.source)
        document = extract_pptx(self.source, self.root / 'table-extract')
        document['translate_all'] = True
        renderer = FakeRenderer()
        snapshots = inspect_layout(self.source, renderer.render(self.source, self.root / 'table-render'),
                                   document, LayoutPolicy(allow_text_compression=True))
        source_style, _ = styled_frame(frame)
        text = source_style.replace('The current value is ', 'Current is ')
        plan = {'reviews':[{'page':1,'needs_changes':True,'issues':['wrap']}],
                'changes':[{'page':1,'shape_id':table.shape_id,'row':1,'column':1,'text':text,'reason':'fit'}]}
        apply_layout_changes(self.source, self.root / 'table-result.pptx', plan, {1:snapshots[1]},
                             LayoutPolicy(allow_text_compression=True), snapshots)
        result = Presentation(self.root / 'table-result.pptx').slides[0].shapes[-1].table.cell(0, 0)
        self.assertEqual(result.text, 'Current is 10 mA')
        self.assertEqual(result.text_frame.paragraphs[0].runs[-1].text, '10 mA')
        self.assertTrue(result.text_frame.paragraphs[0].runs[-1].font.bold)

    def test_wrapping_or_uncertainty_without_visible_damage_never_starts_repair(self):
        class OvercautiousProvider:
            calls = 0
            def generate(inner, prompt, *, images=()):
                inner.calls += 1
                self.assertIn('Screen the whole presentation', prompt)
                return AIResponse(json.dumps({'reviews':[{'page':p,'needs_changes':True,
                    'issues':['one line became two / thumbnail uncertain'],
                    'defects':[{'kind':'wrapping','evidence':'still fits the box'}]}
                    for p in range(1, 13)]}), tuple(images))
        provider = OvercautiousProvider()
        report, renderer = self.modify(provider)
        self.assertEqual(provider.calls, 1)
        self.assertEqual(len(renderer.calls), 2)
        self.assertEqual(report['screening']['selected_pages'], [])

    def test_only_material_overflow_overrides_visual_no_change(self):
        renderer = FakeRenderer()
        before = renderer.render(self.source, self.root / 'metric-reference')
        after = deepcopy(before)
        for page in (2, 3):
            before.slides[page]['text_metrics'][0]['overflow'] = False
            after.slides[page]['text_metrics'][0].update(overflow=True, available_height_pt=100,
                                                       bound_height_pt=102 if page == 2 else 120)
        after.slides[4]['text_metrics'][0]['metric_error'] = True
        candidates = metric_candidates(before, after, list(range(1, 13)))
        self.assertEqual(damage_pages(candidates), [3])

    def test_failed_overview_repairs_only_material_overflow_and_marks_rest_unverified(self):
        class OverflowRenderer(FakeRenderer):
            def render(inner, source, directory, *, pages=None):
                result = super().render(source, directory, pages=pages)
                for page, slide in result.slides.items():
                    for metric in slide['text_metrics']:
                        metric['overflow'] = False
                        if Path(directory).name == 'before' and page == 5:
                            metric.update(overflow=True, available_height_pt=100, bound_height_pt=150)
                return result
        renderer = OverflowRenderer()
        provider = self.provider(reject_screen=True)
        report = PPTXLayoutModifier(provider, renderer, policy=LayoutPolicy(max_passes=1)).modify(
            self.source, self.root / 'overflow-result.pptx', reference=self.source,
            manifest=self.manifest, workdir=self.root / 'overflow-work')
        self.assertEqual(report['screening']['selected_pages'], [5])
        self.assertEqual(renderer.requested_pages[-1], [5])
        self.assertEqual(report['reviewed_pages'], [5])
        self.assertEqual(report['status'], 'needs_review')

    def test_candidate_render_failure_publishes_previous_safe_deck(self):
        class BrokenCandidateRenderer(FakeRenderer):
            def render(inner, source, directory, *, pages=None):
                if Path(directory).name.startswith('after_'):
                    raise KeyError('malformed candidate metrics')
                return super().render(source, directory, pages=pages)
        provider = self.provider(selected=(2,))
        renderer = BrokenCandidateRenderer()
        result = self.root / 'safe-result.pptx'
        report = PPTXLayoutModifier(provider, renderer, policy=LayoutPolicy(max_passes=1)).modify(
            self.source, result, reference=self.source, manifest=self.manifest, workdir=self.root / 'safe-work')
        self.assertEqual(result.read_bytes(), self.source.read_bytes())
        self.assertEqual(report['status'], 'needs_review')
        self.assertEqual(report['failures'][-1]['stage'], 'candidate')
        self.assertEqual(report['failures'][-1]['error_type'], 'KeyError')
        self.assertNotIn(2, report['reviewed_pages'])
