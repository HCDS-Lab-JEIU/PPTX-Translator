import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm

from apply_translation import apply_translation, replace_frame
from gemini_cli import CLIResult
from main import extract_pptx, format_text
from pptx_styles import styled_frame, validate_styles
from pptx_text import equation_roots, frame_text
from lxml import etree
from tests.test_native_math import native_equation
from translate_texts import PPTXTextTranslator, TextRecord, changed_numbers, read_records, write_translation


class SemanticStyleTests(unittest.TestCase):
    def frame(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(Cm(2), Cm(3), Cm(10), Cm(3)).text_frame
        normal = frame.paragraphs[0].add_run()
        normal.text = "Increase "
        normal.font.bold = False
        emphasis = frame.paragraphs[0].add_run()
        emphasis.text = "voltage"
        emphasis.font.bold = True
        emphasis.font.italic = True
        emphasis.font.color.rgb = RGBColor(255, 0, 0)
        return prs, frame

    def test_reordered_translation_carries_bold_color_and_italic_through_saved_pptx(self):
        prs, _ = self.frame()
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, manifest = root / "source.pptx", root / "content.json"
            texts, translated, result = root / "texts.txt", root / "translated.txt", root / "result.pptx"
            prs.save(source)
            before = source.read_bytes()
            document = extract_pptx(source, root / "extract")
            document["translate_all"] = True
            manifest.write_text(json.dumps(document), encoding="utf-8")
            texts.write_text(format_text(document), encoding="utf-8")
            output = "[[STYLE_1]]전압[[/STYLE_1]][[STYLE_0]]을 높이세요[[/STYLE_0]]"
            cli = Mock()
            cli.generate.return_value = CLIResult((), 0, json.dumps({
                "translations": [{"id": 0, "text": output}]}), "")
            PPTXTextTranslator(cli, target_language="Korean").translate_file(texts, translated)
            self.assertIn("[[STYLE_1]]voltage[[/STYLE_1]]", cli.generate.call_args.args[0])
            apply_translation(source, manifest, translated, result)
            frame = Presentation(result).slides[0].shapes[0].text_frame
            self.assertEqual(frame.text, "전압을 높이세요")
            emphasis, normal = frame.paragraphs[0].runs
            self.assertEqual(emphasis.text, "전압")
            self.assertTrue(emphasis.font.bold)
            self.assertTrue(emphasis.font.italic)
            self.assertEqual(emphasis.font.color.rgb, RGBColor(255, 0, 0))
            self.assertFalse(normal.font.bold)
            self.assertEqual(source.read_bytes(), before)
            # Old/checkpoint translations without alignment must not silently use positions.
            write_translation(translated, read_records(texts), {0: "전압을 높이세요"})
            with self.assertRaisesRegex(ValueError, "다시 추출·번역"):
                apply_translation(source, manifest, translated, root / "invalid.pptx")
            self.assertFalse((root / "invalid.pptx").exists())

    def test_missing_duplicate_unknown_nested_and_unmarked_text_rejected(self):
        _, frame = self.frame()
        source, _ = styled_frame(frame)
        for output in (
            "전압을 높이세요",
            source + source,
            source.replace("STYLE_1", "STYLE_9"),
            "[[STYLE_0]][[STYLE_1]]전압[[/STYLE_1]][[/STYLE_0]]",
            source + "추가 텍스트",
        ):
            with self.subTest(output=output), self.assertRaises(ValueError):
                validate_styles(source, output)

    def test_style_ids_are_not_numeric_content(self):
        record = TextRecord(0, 1, "", "Apply 10 V")
        self.assertFalse(changed_numbers(record, "[[STYLE_52]]10 V 적용[[/STYLE_52]]"))
        self.assertTrue(changed_numbers(record, "[[STYLE_10]]20 V 적용[[/STYLE_10]]"))

    def test_invalid_spans_rejected_before_model_output_is_published(self):
        _, frame = self.frame()
        source, _ = styled_frame(frame)
        record = TextRecord(0, 1, "", "Increase voltage", source)
        cli = Mock()
        cli.generate.return_value = CLIResult((), 0, json.dumps({
            "translations": [{"id": 0, "text": "전압을 높이세요"}]}), "")
        with self.assertRaisesRegex(ValueError, "서식 표시"):
            PPTXTextTranslator(cli).translate_batch([record])

    def test_spans_cannot_cross_paragraph_or_equation_boundary(self):
        first = "[[STYLE_0]]Increase[[/STYLE_0]]"
        second = "[[STYLE_1]]voltage[[/STYLE_1]]"
        for boundary in ("\n", "\v", "[[EQ_1]]"):
            with self.subTest(boundary=boundary), self.assertRaises(ValueError):
                validate_styles(first + boundary + second, second + boundary + first)

    def test_soft_breaks_are_preserved(self):
        _, frame = self.frame()
        frame.paragraphs[0].add_line_break()
        frame.paragraphs[0].add_run().text = "10 V"
        source, _ = styled_frame(frame)
        translated = source.replace("Increase ", "Apply ")
        replace_frame(frame, translated)
        self.assertEqual(frame.text, "Apply voltage\v10 V")
        self.assertTrue(frame.paragraphs[0].runs[1].font.bold)

    def test_native_math_xml_in_mixed_style_table_is_preserved(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_table(1, 1, Cm(2), Cm(3), Cm(10), Cm(3)).table.cell(0, 0).text_frame
        paragraph = frame.paragraphs[0]
        normal = paragraph.add_run()
        normal.text = "Potential "
        normal.font.bold = False
        bold = paragraph.add_run()
        bold.text = "energy "
        bold.font.bold = True
        paragraph._p.append(native_equation())
        paragraph.add_run().text = " explanation"
        before = etree.tostring(equation_roots(frame)[0])
        source, _ = styled_frame(frame)
        translated = source.replace("Potential ", "위치 ").replace("energy ", "에너지 ").replace(" explanation", " 설명")
        replace_frame(frame, translated)
        self.assertEqual(frame_text(frame), "위치 에너지 [[EQ_1]] 설명")
        self.assertEqual(etree.tostring(equation_roots(frame)[0]), before)
        self.assertEqual(frame.paragraphs[0].runs[1].text, "에너지 ")
        self.assertTrue(frame.paragraphs[0].runs[1].font.bold)

    def test_adjacent_identical_runs_do_not_split_a_word_into_translation_spans(self):
        _, frame = self.frame()
        paragraph = frame.paragraphs[0]
        paragraph.runs[1].text = "volt"
        rest = paragraph.add_run()
        rest.text = "age"
        # Clone the identical font properties as a spelling/editor split would do.
        from copy import deepcopy
        rest._r.insert(0, deepcopy(paragraph.runs[1]._r.rPr))
        source, _ = styled_frame(frame)
        self.assertIn("[[STYLE_1]]voltage[[/STYLE_1]]", source)
        self.assertNotIn("STYLE_2", source)


if __name__ == "__main__":
    unittest.main()
