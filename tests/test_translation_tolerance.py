"""Formatting changes must not discard otherwise protected translations."""
import json
import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock

from lxml import etree
from pptx import Presentation
from pptx.util import Cm

from apply_translation import apply_translation, replace_frame
from main import extract_pptx, file_sha256, format_text
from pptx_gui.engine import DesktopEngine, RunControl
from pptx_gui.models import AppSettings, FileJob
from pptx_pipeline.ai import AIResponse
from pptx_styles import styled_frame
from pptx_text import equation_roots, frame_text
from tests.test_native_math import native_equation
from translate_texts import PPTXTextTranslator, TextRecord, read_records, write_translation


class TranslationToleranceTests(unittest.TestCase):
    def translator(self, records, texts):
        provider = Mock()
        provider.generate.return_value = AIResponse(json.dumps({'translations':[
            {'id': record.id, 'text': text} for record, text in zip(records, texts)]}))
        translator = PPTXTextTranslator(provider, recover=True)
        output = translator.translate_batch(records)
        return translator, provider, output

    def frame(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(Cm(2), Cm(3), Cm(12), Cm(4)).text_frame
        first = frame.paragraphs[0].add_run()
        first.text = '전압을 '
        first.font.bold = False
        second = frame.paragraphs[0].add_run()
        second.text = '높이세요'
        second.font.bold = True
        return prs, frame

    def test_missing_style_tags_preserve_translation_without_retry(self):
        _, frame = self.frame()
        styled, _ = styled_frame(frame)
        record = TextRecord(0, 1, '', '전압을 높이세요', styled)
        translator, provider, result = self.translator([record], ['Increase voltage'])
        self.assertEqual(result[0], 'Increase voltage')
        self.assertEqual(provider.generate.call_count, 1)
        self.assertEqual(translator.failures, [])
        replace_frame(frame, result[0], recover=True)
        self.assertEqual(frame.text, 'Increase voltage')
        self.assertEqual(len(frame.paragraphs[0].runs), 1)

    def test_empty_particle_span_does_not_destroy_other_semantic_styles(self):
        _, frame = self.frame()
        source, _ = styled_frame(frame)
        output = '[[STYLE_0]][[/STYLE_0]][[STYLE_1]]Increase voltage[[/STYLE_1]]'
        translator, provider, result = self.translator([TextRecord(0, 1, '', frame.text, source)], [output])
        self.assertEqual(translator.failures, [])
        self.assertEqual(provider.generate.call_count, 1)
        replace_frame(frame, result[0], recover=True)
        self.assertEqual(frame.text, 'Increase voltage')
        self.assertTrue(frame.paragraphs[0].runs[0].font.bold)

    def test_unknown_or_broken_style_tags_are_removed_without_losing_prose(self):
        record = TextRecord(0, 1, '', '전압', '[[STYLE_0]]전압[[/STYLE_0]]')
        for text in ('[[STYLE_99]]Voltage[[/STYLE_99]]', '[[STYLE_0]]Voltage[[/STYLE_1]]'):
            with self.subTest(text=text):
                translator, provider, result = self.translator([record], [text])
                self.assertEqual(result[0], 'Voltage')
                self.assertEqual(translator.failures, [])
                self.assertEqual(provider.generate.call_count, 1)

    def test_new_paragraph_count_does_not_leave_literal_style_tags(self):
        _, frame = self.frame()
        source, _ = styled_frame(frame)
        text = '[[STYLE_0]]Increase[[/STYLE_0]]\n[[STYLE_1]]voltage[[/STYLE_1]]'
        translator, _, result = self.translator([TextRecord(0, 1, '', frame.text, source)], [text])
        replace_frame(frame, result[0], recover=True)
        self.assertEqual(frame.text, 'Increase\nvoltage')
        self.assertNotIn('STYLE', frame.text)
        self.assertEqual(translator.failures, [])

    def test_missing_equation_or_changed_number_remains_a_content_failure(self):
        record = TextRecord(0, 1, '', '전압 10 V [[EQ_1]]')
        for text, reason in (('Voltage 20 V [[EQ_1]]', 'numeric_mismatch'),
                             ('Voltage 10 V', 'equation_marker_mismatch')):
            with self.subTest(text=text):
                translator, provider, result = self.translator([record], [text])
                self.assertEqual(result[0], record.text)
                self.assertEqual(provider.generate.call_count, 2)
                self.assertEqual(translator.failures[0]['reason'], reason)
                self.assertIn(reason, provider.generate.call_args.args[0])

    def test_native_equations_keep_xml_when_newlines_and_style_markers_are_lost(self):
        _, frame = self.frame()
        frame.add_paragraph()._p.append(native_equation())
        before = etree.tostring(equation_roots(frame)[0])
        styled, _ = styled_frame(frame)
        record = TextRecord(0, 1, '', frame_text(frame), styled)
        translator, provider, result = self.translator([record], ['Increase voltage [[EQ_1]]'])
        self.assertEqual(provider.generate.call_count, 1)
        self.assertEqual(translator.failures, [])
        replace_frame(frame, result[0], recover=True)
        self.assertEqual(frame_text(frame), 'Increase voltage\n[[EQ_1]]')
        self.assertEqual(etree.tostring(equation_roots(frame)[0]), before)

    def test_application_keeps_plain_mixed_style_translation_and_reports_deferred_formatting(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            prs, _ = self.frame()
            source = root / 'source.pptx'
            prs.save(source)
            original = source.read_bytes()
            document = extract_pptx(source, root / 'extract')
            document['translate_all'] = True
            manifest = root / 'content.json'
            manifest.write_text(json.dumps(document), encoding='utf-8')
            texts = root / 'texts.txt'
            texts.write_text(format_text(document), encoding='utf-8')
            translations = root / 'translated.txt'
            write_translation(translations, read_records(texts), {0: 'Increase voltage'})
            report = apply_translation(source, manifest, translations, root / 'result.pptx', recover=True)
            self.assertEqual(Presentation(root / 'result.pptx').slides[0].shapes[0].text, 'Increase voltage')
            self.assertEqual(report['kept_original'], [])
            self.assertEqual(len(report['formatting_deferred']), 1)
            self.assertEqual(source.read_bytes(), original)

    def test_old_checkpoint_retries_only_failed_items_and_keeps_successes(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            prs = Presentation()
            for text in ('전류', '전압'):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_textbox(Cm(2), Cm(3), Cm(12), Cm(4)).text = text
            source = root / 'source.pptx'
            prs.save(source)
            settings = AppSettings()
            settings.save.directory = str(root / 'results')
            settings.adjust.enabled = False
            work = root / 'work'
            work.mkdir()
            old = {'source_sha256':file_sha256(source),
                'fingerprint':hashlib.sha256(repr(settings).encode()).hexdigest(),
                'translated':{'0':'전류','1':'Voltage'}, 'translation_failed':[1],
                'translation_failures':[{'page':1,'record_id':0,'reason':'old style failure'}],
                'applied':False,'adjusted_pages':[],'review_reports':[]}
            (work / 'checkpoint.json').write_text(json.dumps(old), encoding='utf-8')
            job = FileJob(source)
            job.workdir = str(work)
            provider = Mock()
            provider.generate.return_value = AIResponse('{"translations":[{"id":0,"text":"Current"}]}')
            engine = DesktopEngine(settings, {}, lambda event: None, RunControl())
            engine.provider = lambda *args, **kwargs: provider
            result = engine.run_file(job)
            requested = json.loads(provider.generate.call_args.args[0].split('Input records:\n')[1])
            self.assertEqual([r['id'] for r in requested], [0])
            self.assertEqual(provider.generate.call_count, 1)
            self.assertEqual(result['translation_failed_pages'], [])
            self.assertEqual([s.shapes[0].text for s in Presentation(result['output']).slides], ['Current','Voltage'])
            state = json.loads((work / 'checkpoint.json').read_text(encoding='utf-8'))
            self.assertEqual(state['prior_translation_failures'], old['translation_failures'])
