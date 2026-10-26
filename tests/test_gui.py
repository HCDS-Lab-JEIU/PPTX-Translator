"""Desktop behavior and page checkpoint tests. Qt cases need requirements-gui.txt."""
from copy import deepcopy
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from pptx import Presentation
from pptx.util import Cm
from pptx_gui.models import AppSettings,SaveSettings,FileJob,Connection,ModelInfo,QuotaBucket,validate_model
from pptx_gui.adapters.codex import parse_limits
from pptx_gui.adapters.base import safe_error
from pptx_gui.engine import DesktopEngine,RunControl,Stopped,QuotaPause
from pptx_pipeline.ai import AIResponse
from main import file_sha256

os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
try:
    from PySide6.QtCore import QSettings,QTimer
    from PySide6.QtWidgets import QApplication
    from PySide6.QtTest import QTest
    from pptx_gui.main_window import MainWindow
    HAVE_QT=True
except ImportError:
    HAVE_QT=False


class DesktopDataTests(unittest.TestCase):
    def test_save_name_and_original_protection(self):
        with TemporaryDirectory() as tmp:
            src=Path(tmp)/'Lecture.pptx'
            cfg=SaveSettings(tmp)
            self.assertEqual(cfg.target(src,'English').name,'Lecture_영어.pptx')
            cfg.filename='{stem}.pptx'
            with self.assertRaises(ValueError): cfg.target(src,'English')
            cfg.filename='../escape.pptx'
            with self.assertRaises(ValueError): cfg.target(src,'English')

    def test_null_limits_remain_unknown_and_all_windows_are_preserved(self):
        buckets=parse_limits({'rateLimitsByLimitId':{'a':{'primary':{'usedPercent':None,'windowDurationMins':45},
            'secondary':{'usedPercent':25,'resetsAt':123}},'b':{'primary':{'usedPercent':100}}}})
        self.assertEqual([b.remaining for b in buckets],[None,75,0])
        self.assertEqual(buckets[0].window_minutes,45)
        self.assertEqual(buckets[1].resets_at,123)

    def test_invalid_model_and_effort_are_revalidated(self):
        cfg=AppSettings().translate; cfg.model='not-available'; cfg.effort='max'
        connection=Connection('codex',models=[ModelInfo('valid','Valid',('low','medium'))])
        validate_model(cfg,connection)
        self.assertEqual((cfg.model,cfg.effort),('valid','low'))
        validate_model(cfg,Connection('gemini'))
        self.assertEqual((cfg.model,cfg.effort),('',''))

    def test_credentials_are_redacted(self):
        message='{"access_token":"private-token","api_key":"private-key"} Bearer bearer-value https://host/callback?code=private-code'
        clean=safe_error(message)
        for secret in ('private-token','private-key','bearer-value','private-code'): self.assertNotIn(secret,clean)

    def test_quota_pause_switches_connection_without_changing_current_language(self):
        cfg=AppSettings(); control=RunControl(); events=[]
        engine=DesktopEngine(cfg,{},events.append,control)
        engine.current_settings=deepcopy(cfg)
        pending=deepcopy(cfg); pending.translate.provider='gemini'; pending.translate.target_language='Japanese'
        calls=[]
        def emit(event):
            events.append(event)
            if event['type']=='quota_paused':
                engine.update_settings(pending,quota_resume=True)
                control.resume_event.set()
        engine.emit=emit
        def unit():
            calls.append(1)
            if len(calls)==1: raise QuotaPause()
            return 'ok'
        self.assertEqual(engine.unit(unit),'ok')
        self.assertEqual(engine.current_settings.translate.provider,'gemini')
        self.assertEqual(engine.current_settings.translate.target_language,'English')
        self.assertEqual(engine.settings.translate.target_language,'Japanese')

    def test_size_split_checkpoint_resume_and_english_to_korean_application(self):
        with TemporaryDirectory() as tmp:
            folder=Path(tmp); source=folder/'source.pptx'
            prs=Presentation()
            for text in ('Current','Voltage'):
                slide=prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_textbox(Cm(2),Cm(3),Cm(10),Cm(3)).text=text
            prs.save(source); original=file_sha256(source)
            cfg=AppSettings(); cfg.save.directory=str(folder/'results'); cfg.adjust.enabled=False
            cfg.translate.source_language='English'; cfg.translate.target_language='Korean'
            job=FileJob(source)
            class Fake:
                model='test'
                calls=0
                fail=True
                def generate(inner,prompt,**kwargs):
                    data=json.loads(prompt.split('Input records:\n')[1])
                    inner.calls+=1
                    if inner.calls==2 and inner.fail: raise Stopped()
                    return AIResponse(json.dumps({'translations':[{'id':r['id'],'text':{'Current':'전류','Voltage':'전압'}[r['text']]} for r in data]}))
            provider=Fake(); engine=DesktopEngine(cfg,{},lambda e:None,RunControl())
            engine.provider=lambda node,layout=False:provider
            # Force multiple requests only to exercise interruption between batches.
            with patch('pptx_gui.engine.DEFAULT_BATCH_CHARS', 1):
                with self.assertRaises(Stopped): engine.run_file(job)
            checkpoint=json.loads((Path(job.workdir)/'checkpoint.json').read_text(encoding='utf-8'))
            self.assertEqual(len(checkpoint['translated']),1)
            provider.fail=False; result=engine.run_file(job)
            self.assertEqual(provider.calls,3)
            result_prs=Presentation(result['output'])
            self.assertEqual([s.shapes[0].text for s in result_prs.slides],['전류','전압'])
            self.assertEqual(file_sha256(source),original)
            with self.assertRaises(FileExistsError): engine.run_file(job)


@unittest.skipUnless(HAVE_QT,'Install requirements-gui.txt for Qt UI tests')
class DesktopWidgetTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): cls.app=QApplication.instance() or QApplication([])
    def setUp(self):
        self.temp=TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        prefs=QSettings(str(Path(self.temp.name)/'ui.ini'),QSettings.Format.IniFormat)
        self.window=MainWindow(demo=True,preferences=prefs)
        self.addCleanup(self.close)
    def close(self):
        if self.window.running():
            self.window.stop()
            for _ in range(50):
                QTest.qWait(30)
                if not self.window.running(): break
        self.window.close(); self.app.processEvents()
    def test_node_settings_are_independent_and_shared_quota_is_account_scoped(self):
        self.window.settings.translate.prompt='Translation instruction'
        self.window.select_node('adjust')
        self.window.settings.adjust.prompt='Layout instruction'
        self.window.settings.adjust.effort='high'
        self.window.select_node('translate')
        self.assertEqual(self.window.settings.translate.prompt,'Translation instruction')
        self.assertEqual(self.window.settings.translate.effort,'medium')
        self.assertEqual(self.window.settings.adjust.effort,'high')
        a=self.window.connections.state(self.window.settings.translate.provider)
        b=self.window.connections.state(self.window.settings.adjust.provider)
        self.assertIs(a,b)
        self.window.inspector.provider_changed('gemini')
        self.assertEqual(self.window.settings.translate.effort,'')
        self.assertEqual(self.window.settings.adjust.provider,'codex')
    def test_save_node_uses_selected_file_for_preview(self):
        self.window.inspector.source=Path('example.pptx')
        self.window.select_node('save')
        from PySide6.QtWidgets import QLabel
        self.assertTrue(any('example_영어.pptx' in label.text() for label in self.window.inspector.findChildren(QLabel)))
    def test_antigravity_can_be_selected_for_translation_and_layout(self):
        self.window.inspector.provider_changed('antigravity')
        self.assertEqual(self.window.settings.translate.model,'demo-agy')
        self.window.select_node('adjust'); self.window.inspector.provider_changed('antigravity')
        self.assertEqual(self.window.settings.adjust.model,'demo-agy')
    def test_adjustment_compression_can_be_disabled_and_survives_node_switch(self):
        from PySide6.QtWidgets import QCheckBox
        self.window.select_node('adjust')
        checkbox = next(w for w in self.window.inspector.findChildren(QCheckBox)
                        if w.text() == '필요한 문장만 의미를 유지하며 압축')
        self.assertTrue(checkbox.isChecked())
        checkbox.setChecked(False)
        self.window.select_node('translate')
        self.window.select_node('adjust')
        self.assertFalse(self.window.settings.adjust.allow_text_compression)
    def test_account_switch_does_not_reuse_previous_account_quota(self):
        manager=self.window.connections
        manager.states['codex']=Connection('codex',connected=True,account='account A',buckets=[QuotaBucket('A',75)])
        manager._result('codex',Connection('codex',connected=True,account='account B',quota_status='조회 불가'))
        self.assertEqual(manager.state('codex').buckets,[])
    def test_pause_is_responsive_and_stop_keeps_ui_alive(self):
        beats=[]; timer=QTimer(); timer.setInterval(20); timer.timeout.connect(lambda:beats.append(1)); timer.start()
        self.window.start(); QTest.qWait(200); self.window.pause_resume()
        for _ in range(30):
            QTest.qWait(30)
            if self.window.paused: break
        self.assertTrue(self.window.paused)
        before=len(beats); QTest.qWait(120); self.assertGreater(len(beats),before)
        self.window.select_node('save')
        self.window.stop()
        for _ in range(50):
            QTest.qWait(30)
            if not self.window.running(): break
        self.assertFalse(self.window.running()); timer.stop()


if __name__=='__main__': unittest.main()
