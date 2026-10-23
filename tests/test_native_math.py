from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.util import Cm

import main
from apply_translation import apply_translation, replace_frame
from pptx_text import A, M, MC, frame_text, equation_roots, math_text_nodes
from translate_texts import TextRecord, changed_equations, parse_records, write_translation


def native_equation():
    return etree.fromstring(('''<a14:m xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"
        xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
        <m:oMath><m:r><m:t>v=</m:t></m:r><m:f><m:num><m:r><m:t>dW</m:t></m:r></m:num>
        <m:den><m:r><m:t>dq</m:t></m:r></m:den></m:f></m:oMath></a14:m>''').encode())


class NativeMathTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.source = self.folder / "source.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Cm(2), Cm(3), Cm(10), Cm(4))
        shape.text_frame.paragraphs[0].text = "전하의 위치 에너지 "
        shape.text_frame.paragraphs[0]._p.append(native_equation())
        shape.text_frame.paragraphs[0].add_run().text = " 설명"
        shape.text_frame.add_paragraph().text = "전압"
        fallback = deepcopy(shape.element)
        for text in fallback.iter("{" + A + "}t"):
            text.text = "Legacy fallback"
        wrapper = etree.Element("{" + MC + "}AlternateContent", nsmap={"mc": MC})
        choice = etree.SubElement(wrapper, "{" + MC + "}Choice", Requires="a14",
                                  nsmap={"a14": "http://schemas.microsoft.com/office/drawing/2010/main"})
        slide.shapes._spTree.replace(shape.element, wrapper)
        choice.append(shape.element)
        etree.SubElement(wrapper, "{" + MC + "}Fallback").append(fallback)
        # A native table under another AlternateContent, with Korean inside Math.
        table_shape = slide.shapes.add_table(1, 1, Cm(2), Cm(9), Cm(10), Cm(3))
        cell = table_shape.table.cell(0, 0)
        cell.text = "전류 "
        equation = native_equation()
        next(equation.iter("{" + M + "}t")).text = "전하량"
        cell.text_frame.paragraphs[0]._p.append(equation)
        wrapper2 = etree.Element("{" + MC + "}AlternateContent", nsmap={"mc": MC})
        choice2 = etree.SubElement(wrapper2, "{" + MC + "}Choice", Requires="a14",
                                   nsmap={"a14": "http://schemas.microsoft.com/office/drawing/2010/main"})
        slide.shapes._spTree.replace(table_shape.element, wrapper2)
        choice2.append(table_shape.element)
        prs.save(self.source)

    def extract(self):
        return main.extract_pptx(self.source, self.folder / "output")

    def test_extracts_choice_once_and_protects_equations_in_shapes_and_tables(self):
        document = self.extract()
        records = document["slides"][0]["texts"]
        self.assertEqual(len(records), 3)
        self.assertEqual({r["kind"] for r in records}, {"text", "table_cell", "math_text"})
        self.assertIn("[[EQ_1]]", records[0]["text"])
        self.assertNotIn("Legacy fallback", main.format_text(document))
        label = next(r for r in records if r["kind"] == "math_text")
        self.assertEqual(label["text"], "전하량")
        self.assertEqual(label["row"], 1)
        self.assertIn("math_text=1", main.format_text(document))

    def test_translates_native_prose_and_labels_without_changing_formula_structure(self):
        document = self.extract()
        manifest = self.folder / "content.json"
        manifest.write_text(json.dumps(document), encoding="utf-8")
        records = parse_records(main.format_text(document))
        translated = self.folder / "translated.txt"
        values = {}
        for record in records:
            if "전하의" in record.text:
                values[record.id] = "Potential energy [[EQ_1]] description\nVoltage"
            elif record.text == "전하량":
                values[record.id] = "Charge"
            else:
                values[record.id] = "Current [[EQ_1]]"
        write_translation(translated, records, values)
        output = self.folder / "translated.pptx"
        report = apply_translation(self.source, manifest, translated, output)
        self.assertEqual(report["replaced"], 3)
        source_prs, translated_prs = Presentation(self.source), Presentation(output)
        source_shapes = [s for s, _ in main.walk_shapes(source_prs.slides[0].shapes)]
        translated_shapes = [s for s, _ in main.walk_shapes(translated_prs.slides[0].shapes)]
        self.assertEqual(frame_text(translated_shapes[0].text_frame), values[0])
        self.assertEqual(etree.tostring(equation_roots(source_shapes[0].text_frame)[0]),
                         etree.tostring(equation_roots(translated_shapes[0].text_frame)[0]))
        old_math = deepcopy(equation_roots(source_shapes[1].table.cell(0, 0).text_frame)[0])
        new_math = deepcopy(equation_roots(translated_shapes[1].table.cell(0, 0).text_frame)[0])
        next(old_math.iter("{" + M + "}t")).text = "Charge"
        self.assertEqual(etree.tostring(old_math), etree.tostring(new_math))
        with ZipFile(self.source) as old, ZipFile(output) as new:
            for name in old.namelist():
                if name != "ppt/slides/slide1.xml":
                    self.assertEqual(old.read(name), new.read(name))

    def test_missing_or_moved_equation_marker_is_rejected(self):
        prs = Presentation(self.source)
        shape = next(main.walk_shapes(prs.slides[0].shapes))[0]
        for invalid in ("Energy\nVoltage", "Energy\nVoltage [[EQ_1]]"):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                replace_frame(shape.text_frame, invalid)
        original = TextRecord(0, 1, "", "에너지 [[EQ_1]]\n전압")
        self.assertTrue(changed_equations(original, "Energy\nVoltage [[EQ_1]]"))
        self.assertTrue(changed_equations(original, "Energy\nVoltage"))
        self.assertFalse(changed_equations(original, "Energy [[EQ_1]]\nVoltage"))

    def test_legacy_extraction_cannot_silently_skip_native_math(self):
        document = self.extract()
        document.pop("text_schema_version")
        manifest = self.folder / "content.json"
        manifest.write_text(json.dumps(document), encoding="utf-8")
        records = parse_records(main.format_text(document))
        translated = self.folder / "translated.txt"
        write_translation(translated, records, {r.id: r.text for r in records})
        with self.assertRaisesRegex(ValueError, "수식 지원 전"):
            apply_translation(self.source, manifest, translated)


if __name__ == "__main__":
    unittest.main()
