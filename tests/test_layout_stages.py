"""Ordered repair strategies, verified candidate removal and content protection."""
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation
from pptx.util import Cm, Pt

from main import extract_pptx
from pptx_pipeline.ai import AIResponse
from pptx_pipeline.layout import LayoutPolicy, PPTXLayoutModifier, inspect_layout, apply_layout_changes
from pptx_pipeline.layout_adjustments import STAGES, apply_whitespace, validate_whitespace
from tests.test_pipeline import FakeRenderer, PIXEL


class StagedLayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.pptx'
        prs = Presentation()
        for _ in range(4):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            text = slide.shapes.add_textbox(Cm(2), Cm(3), Cm(12), Cm(3))
            text.text = 'Current value is 10 mA'
            text.text_frame.paragraphs[0].font.size = Pt(20)
            slide.shapes.add_picture(BytesIO(PIXEL), Cm(16), Cm(10), Cm(3), Cm(2))
        prs.save(self.source)
        self.document = extract_pptx(self.source, self.root / 'extract')
        self.document['translate_all'] = True
        self.manifest = self.root / 'content.json'
        self.manifest.write_text(json.dumps(self.document), encoding='utf-8')

    def provider(self, solve_all_with_font=False):
        class Provider:
            model = 'simulated-stages'
            def __init__(inner): inner.requests = []
            def generate(inner, prompt, *, images=()):
                if 'Screen the whole presentation' in prompt:
                    return AIResponse(json.dumps({'reviews':[{'page':p,'needs_changes':True,'issues':['overlap'],
                        'defects':[{'kind':'overlap','evidence':'New text obscures a diagram label'}],
                        'predicted_adjustments':list(STAGES)}
                        for p in range(1,5)]}), tuple(images))
                data = json.loads(prompt[prompt.index('{"images":'):])
                stage, verification = data['stage'], data['verification_only']
                pages = [slide['page'] for slide in data['slides']]
                inner.requests.append((stage, verification, pages))
                reviews, changes = [], []
                for page in pages:
                    target = 'font' if solve_all_with_font else STAGES[page-1]
                    solves = target == stage
                    reviews.append({'page':page,'needs_changes':not (verification and solves), 'issues':[]})
                    if solves and not verification:
                        action = {'font':{'font_scale':.9}, 'spacing':{'spacing_scale':.85},
                            'images':{'box_cm':{'x':17,'y':10,'width':2.4,'height':1.6}},
                            'compression':{'text':'Current is 10 mA'}}[stage]
                        changes.append({'page':page,'shape_id':3 if stage=='images' else 2,'reason':'fit',**action})
                return AIResponse(json.dumps({'reviews':reviews,'changes':changes}), tuple(images))
        return Provider()

    def run_modifier(self, provider, compression=True):
        renderer = FakeRenderer()
        report = PPTXLayoutModifier(provider, renderer, staged=True,
            policy=LayoutPolicy(max_passes=1, allow_text_compression=compression)).modify(
                self.source, self.root / 'result.pptx', reference=self.source,
                manifest=self.manifest, workdir=self.root / 'work')
        return report, renderer

    def test_candidates_leave_pipeline_only_after_rendered_verification(self):
        original = self.source.read_bytes()
        provider = self.provider()
        report, renderer = self.run_modifier(provider)
        self.assertEqual(provider.requests, [
            ('font',False,[1,2,3,4]), ('font',True,[1]),
            ('spacing',False,[2,3,4]), ('spacing',True,[2]),
            ('images',False,[3,4]), ('images',True,[3]),
            ('compression',False,[4]), ('compression',True,[4])])
        self.assertEqual([stage['resolved_pages'] for stage in report['pipeline_stages']], [[1],[2],[3],[4]])
        self.assertEqual([stage['remaining_pages'] for stage in report['pipeline_stages']], [[2,3,4],[3,4],[4],[]])
        self.assertEqual(renderer.requested_pages[2:], [[1],[2],[3],[4]])
        self.assertEqual(report['unresolved'], [])
        self.assertEqual(report['status'], 'needs_review')  # compressed prose stays reviewable
        result = Presentation(self.root / 'result.pptx')
        self.assertEqual(result.slides[0].shapes[0].text_frame.paragraphs[0].font.size, Pt(18))
        self.assertEqual(result.slides[3].shapes[0].text, 'Current is 10 mA')
        self.assertEqual(self.source.read_bytes(), original)

    def test_font_solution_skips_all_later_strategies(self):
        provider = self.provider(solve_all_with_font=True)
        report, renderer = self.run_modifier(provider)
        self.assertEqual([stage['stage'] for stage in report['pipeline_stages']], ['font'])
        self.assertEqual(provider.requests, [('font',False,[1,2,3,4]), ('font',True,[1,2,3,4])])
        self.assertEqual(report['status'], 'verified')
        self.assertEqual(len(renderer.calls), 3)

    def test_disabled_compression_is_never_requested(self):
        provider = self.provider()
        report, _ = self.run_modifier(provider, compression=False)
        self.assertEqual(report['stage_order'], ['font','spacing','images'])
        self.assertFalse(any(stage=='compression' for stage,_,_ in provider.requests))
        self.assertEqual([review['page'] for review in report['unresolved']], [4])

    def test_resume_reuses_stage_specific_reviews_and_renders(self):
        provider = self.provider()
        self.run_modifier(provider)
        calls = list(provider.requests)
        report, renderer = self.run_modifier(provider)
        self.assertEqual(provider.requests, calls)
        self.assertEqual(renderer.calls, [])
        self.assertEqual([stage['resolved_pages'] for stage in report['pipeline_stages']], [[1],[2],[3],[4]])

    def test_failed_verification_does_not_remove_candidates(self):
        provider = self.provider(solve_all_with_font=True)
        original = provider.generate
        def generate(prompt, *, images=()):
            if '"verification_only": true' in prompt:
                return AIResponse('{}', tuple(images))
            return original(prompt, images=images)
        provider.generate = generate
        report, _ = self.run_modifier(provider, compression=False)
        self.assertEqual(report['pipeline_stages'][0]['resolved_pages'], [])
        self.assertEqual(report['pipeline_stages'][1]['input_pages'], [1,2,3,4])
        self.assertEqual([r['page'] for r in report['unresolved']], [1,2,3,4])

    def test_blank_removal_includes_leading_and_trailing_paragraphs(self):
        frame = Presentation(self.source).slides[0].shapes[0].text_frame
        frame.text = '\n\nCurrent   value\n\n'
        frame.paragraphs[2].level = 2
        target = 'Current value'
        validate_whitespace({'text':frame.text,'text_editable':True}, target)
        apply_whitespace(frame, target)
        self.assertEqual(frame.text, target)
        self.assertEqual(frame.paragraphs[0].level, 2)

    def test_spacing_cumulative_limit_and_equation_line_boundaries(self):
        renderer, policy = FakeRenderer(), LayoutPolicy()
        snapshots = inspect_layout(self.source, renderer.render(self.source, self.root / 'render'), self.document, policy)
        plan = {'reviews':[{'page':1,'needs_changes':True,'issues':[]}],
                'changes':[{'page':1,'shape_id':2,'reason':'fit','spacing_scale':.7}]}
        first = self.root / 'first.pptx'
        apply_layout_changes(self.source, first, plan, {1:snapshots[1]}, policy, snapshots, stage='spacing')
        current = inspect_layout(first, renderer.render(first, self.root / 'after'), self.document, policy, snapshots)
        with self.assertRaisesRegex(ValueError, '누적'):
            apply_layout_changes(first, self.root / 'second.pptx', plan, {1:current[1]}, policy, snapshots, stage='spacing')
        with self.assertRaises(ValueError):
            validate_whitespace({'text':'Voltage\n\n[[EQ_1]]','text_editable':True}, 'Voltage\n[[EQ_1]]')

    def test_stage_boundaries_reject_early_rewrite_and_textbox_movement(self):
        renderer = FakeRenderer()
        policy = LayoutPolicy(allow_text_compression=True)
        snapshots = inspect_layout(self.source, renderer.render(self.source, self.root / 'render'), self.document, policy)
        for stage, change in (('font',{'text':'Current is 10 mA'}), ('spacing',{'font_scale':.9}),
                              ('images',{'box_cm':{'x':3,'y':3,'width':12,'height':3}})):
            with self.subTest(stage=stage):
                plan = {'reviews':[{'page':1,'needs_changes':True,'issues':[]}],
                        'changes':[{'page':1,'shape_id':2,'reason':'test',**change}]}
                output = self.root / f'{stage}.pptx'
                with self.assertRaises(ValueError):
                    apply_layout_changes(self.source, output, plan, {1:snapshots[1]}, policy, snapshots, stage=stage)
                self.assertFalse(output.exists())

    def test_literal_whitespace_preserves_bold_and_paragraph_properties(self):
        prs = Presentation(self.source)
        frame = prs.slides[0].shapes[0].text_frame
        frame.clear()
        first = frame.paragraphs[0].add_run()
        first.text = 'Increase   '
        bold = frame.paragraphs[0].add_run()
        bold.text = 'voltage'
        bold.font.bold = True
        frame.add_paragraph()
        frame.add_paragraph()
        last = frame.add_paragraph()
        last.text = 'Use  10 mA'
        last.level = 1
        target = 'Increase voltage\nUse 10 mA'
        validate_whitespace({'text':frame.text,'text_editable':True}, target)
        apply_whitespace(frame, target)
        self.assertEqual(frame.text, target)
        self.assertTrue(frame.paragraphs[0].runs[-1].font.bold)
        self.assertEqual(frame.paragraphs[1].level, 1)

    def test_whitespace_cannot_join_words_change_numbers_or_code_indentation(self):
        for original,target in (('do not','donot'), ('10 mA','10mA'), ('10 mA','20 mA'),
                                ('return   current','return current')):
            with self.subTest(original=original), self.assertRaises(ValueError):
                validate_whitespace({'text':original,'text_editable':True}, target)
