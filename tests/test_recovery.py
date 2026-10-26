"""Recoverable failures preserve safe work, XML and review status without AI calls."""
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

from pptx import Presentation
from pptx.util import Cm

from apply_translation import apply_translation, replace_frame
from gemini_cli import EmptyModelResponseError
from main import extract_pptx, file_sha256, format_text
from pptx_gui.engine import DesktopEngine, RunControl, Stopped
from pptx_gui.models import AppSettings, FileJob
from pptx_pipeline.ai import AIResponse
from translate_texts import DEFAULT_BATCH_CHARS, PPTXTextTranslator, TextRecord, make_batches, read_records, write_translation


def response(items):
    return AIResponse(json.dumps({'translations': items}))


class RecoveryTests(unittest.TestCase):
    def test_retry_only_invalid_records_and_keep_good_translation(self):
        records = [TextRecord(0, 1, '', '전류'), TextRecord(1, 1, '', '10 V'),
                   TextRecord(2, 1, '', '전압', '[[STYLE_0]]전[[/STYLE_0]][[STYLE_1]]압[[/STYLE_1]]')]
        cli = Mock()
        cli.generate.side_effect = [response([{'id': 0, 'text': 'Current'},
            {'id': 1, 'text': '20 V'}, {'id': 2, 'text': 'Voltage'}]),
            response([{'id': 1, 'text': '10 volts'}, {'id': 2, 'text': 'Voltage'}])]
        translator = PPTXTextTranslator(cli, recover=True)
        self.assertEqual(translator.translate_batch(records), {0: 'Current', 1: '10 volts', 2: 'Voltage'})
        retried = json.loads(cli.generate.call_args.args[0].split('Input records:\n')[1].split('\nRetry:')[0])
        self.assertEqual([r['id'] for r in retried], [1])
        self.assertEqual(translator.failures, [])
        self.assertEqual([r['record_id'] for r in translator.deferred], [2])

    def test_malformed_missing_duplicate_and_empty_responses_keep_original(self):
        record = TextRecord(0, 1, '', '전압 10 V')
        for bad in (AIResponse('invalid json'), response([]),
                    response([{'id': 0, 'text': '10 V'}] * 2), EmptyModelResponseError('empty')):
            with self.subTest(bad=type(bad).__name__):
                cli = Mock()
                if isinstance(bad, Exception):
                    cli.generate.side_effect = bad
                else:
                    cli.generate.return_value = bad
                translator = PPTXTextTranslator(cli, recover=True)
                self.assertEqual(translator.translate_batch([record]), {0: record.text})
                self.assertEqual(cli.generate.call_count, 2)
                self.assertEqual(len(translator.failures), 1)

    def test_equation_line_break_is_restored_without_discarding_translation(self):
        record = TextRecord(0, 1, '', '전압\n[[EQ_1]]')
        cli = Mock()
        cli.generate.return_value = response([{'id': 0, 'text': 'Voltage [[EQ_1]]'}])
        translator = PPTXTextTranslator(cli, recover=True)
        self.assertEqual(translator.translate_batch([record]), {0: 'Voltage\n[[EQ_1]]'})
        self.assertEqual(translator.failures, [])
        self.assertEqual(cli.generate.call_count, 1)
        self.assertIn('equation_lines_restored', translator.deferred[0]['reasons'])

    def make_source(self, folder):
        source = folder / 'source.pptx'
        prs = Presentation()
        for text in ('전류', '전압'):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Cm(2), Cm(3), Cm(10), Cm(3)).text = text
        prs.save(source)
        return source

    def test_application_restores_partial_mutation_and_continues(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = self.make_source(root)
            original = file_sha256(source)
            document = extract_pptx(source, root / 'extract')
            manifest = root / 'content.json'
            manifest.write_text(json.dumps(document), encoding='utf-8')
            texts = root / 'texts.txt'
            texts.write_text(format_text(document), encoding='utf-8')
            translated = root / 'translated.txt'
            write_translation(translated, read_records(texts), {0: 'Current', 1: 'Voltage'})
            def fail_after_mutation(frame, text, **kwargs):
                if text == 'Current':
                    frame.text = 'CORRUPTED'
                    raise ValueError('simulated post-mutation verification failure')
                return replace_frame(frame, text, **kwargs)
            with patch('apply_translation.replace_frame', side_effect=fail_after_mutation):
                report = apply_translation(source, manifest, translated, root / 'result.pptx', recover=True)
            result = Presentation(root / 'result.pptx')
            self.assertEqual([s.shapes[0].text for s in result.slides], ['전류', 'Voltage'])
            self.assertEqual(report['kept_original'][0]['page'], 1)
            self.assertEqual(report['replaced'], 1)
            self.assertEqual(file_sha256(source), original)

    def test_engine_saves_and_marks_review_after_translation_and_layout_failures(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = self.make_source(root)
            original = file_sha256(source)
            cfg = AppSettings()
            cfg.save.directory = str(root / 'results')
            events = []
            engine = DesktopEngine(cfg, {}, events.append, RunControl())
            cli = Mock()
            cli.generate.side_effect = [response([{'id': 0, 'text': ''}, {'id': 1, 'text': 'Voltage'}]),
                                       response([{'id': 0, 'text': ''}])]
            engine.provider = lambda *args, **kwargs: cli
            job = FileJob(source)
            def unsafe_candidate(source, destination, **kwargs):
                Path(destination).write_bytes(b'unverified candidate')
                raise RuntimeError('render failed')
            with patch('pptx_gui.engine.PPTXLayoutModifier.modify', side_effect=unsafe_candidate) as modify:
                result = engine.run_file(job)
            self.assertEqual(modify.call_count, 1)
            self.assertEqual(modify.call_args.kwargs['pages'], [1, 2])
            self.assertEqual(result['status'], 'needs_review')
            self.assertEqual(result['translation_failed_pages'], [1])
            output = Presentation(result['output'])
            self.assertEqual([s.shapes[0].text for s in output.slides], ['전류', 'Voltage'])
            self.assertEqual(file_sha256(source), original)
            self.assertEqual(len(result['review_reports']), 2)
            self.assertTrue(Path(result['output']).with_suffix('.report.json').exists())
            self.assertEqual(sum(e['type'] == 'warning' for e in events), 2)

    def test_engine_does_not_swallow_stop_or_resource_exhaustion_during_adjustment(self):
        for error in (Stopped(), MemoryError(), RecursionError()):
            with self.subTest(error=type(error).__name__), TemporaryDirectory() as tmp:
                root = Path(tmp)
                source = self.make_source(root)
                cfg = AppSettings()
                cfg.save.directory = str(root / 'results')
                cfg.translate.enabled = False
                engine = DesktopEngine(cfg, {}, lambda e: None, RunControl())
                engine.provider = lambda *args, **kwargs: Mock()
                with patch('pptx_gui.engine.PPTXLayoutModifier.modify', side_effect=error):
                    with self.assertRaises(type(error)):
                        engine.run_file(FileJob(source))
                self.assertFalse(cfg.save.target(source, cfg.translate.target_language).exists())

    def test_optional_adjustment_exceptions_still_save_the_valid_translation(self):
        for error in (PermissionError('work file locked'), KeyError('metrics'), AttributeError('missing metric'),
                      TypeError('invalid plan'), IndexError('missing frame'), StopIteration(), RuntimeError('render'),
                      ModuleNotFoundError("No module named 'pypdfium2'", name='pypdfium2')):
            with self.subTest(error=type(error).__name__), TemporaryDirectory() as tmp:
                root = Path(tmp)
                source = self.make_source(root)
                cfg = AppSettings()
                cfg.save.directory = str(root / 'results')
                cfg.translate.enabled = False
                events = []
                engine = DesktopEngine(cfg, {}, events.append, RunControl())
                engine.provider = lambda *args, **kwargs: Mock()
                def fail_adjustment(*args, **kwargs):
                    kwargs['progress']('AI 시각 검토 1: 페이지 [1, 2]')
                    raise error
                with patch('pptx_gui.engine.PPTXLayoutModifier.modify', side_effect=fail_adjustment):
                    result = engine.run_file(FileJob(source))
                self.assertEqual(result['status'], 'needs_review')
                self.assertEqual(Path(result['output']).read_bytes(), source.read_bytes())
                self.assertTrue(Path(result['output']).with_suffix('.report.json').is_file())
                self.assertEqual(result['review_reports'][0]['unresolved'][0]['error_type'], type(error).__name__)
                self.assertEqual(result['review_reports'][0]['unresolved'][0]['stage'], 'AI 시각 검토 1: 페이지 [1, 2]')
                self.assertTrue(any(type(error).__name__ in event.get('message', '') for event in events))
                if isinstance(error, ModuleNotFoundError):
                    self.assertIn('pypdfium2', result['review_reports'][0]['unresolved'][0]['detail'])
                    self.assertTrue(any('pypdfium2' in event.get('message', '') for event in events))

    def test_checkpoint_retries_transient_lock_but_rejects_persistent_denial(self):
        for persistent in (False, True):
            with self.subTest(persistent=persistent), TemporaryDirectory() as tmp:
                root = Path(tmp)
                source = self.make_source(root)
                cfg = AppSettings()
                cfg.save.directory = str(root / 'results')
                cfg.translate.enabled = cfg.adjust.enabled = False
                engine = DesktopEngine(cfg, {}, lambda e: None, RunControl())
                original_replace = Path.replace
                failures = []
                def locked_replace(path, destination):
                    if path.name == 'checkpoint.tmp' and (persistent or not failures):
                        failures.append(1)
                        raise PermissionError('simulated Windows lock')
                    return original_replace(path, destination)
                with patch.object(Path, 'replace', locked_replace), patch('pptx_gui.engine.time.sleep'):
                    if persistent:
                        with self.assertRaises(PermissionError):
                            engine.run_file(FileJob(source))
                        self.assertEqual(len(failures), 6)
                    else:
                        result = engine.run_file(FileJob(source))
                        self.assertEqual(result['status'], 'done')
                        self.assertEqual(len(failures), 1)

    def test_entire_47_slide_deck_uses_one_translation_request(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / 'source.pptx'
            prs = Presentation()
            for index in range(47):
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_textbox(Cm(2), Cm(3), Cm(10), Cm(3)).text = f'전류 {index}'
            prs.save(source)
            cfg = AppSettings()
            cfg.save.directory = str(root / 'results')
            cfg.adjust.enabled = False
            events = []
            engine = DesktopEngine(cfg, {}, events.append, RunControl())
            requests = []
            def translate(prompt, **kwargs):
                data = json.loads(prompt.split('Input records:\n')[1])
                requests.append(data)
                return response([{'id': r['id'], 'text': r['text'].replace('전류', 'Current')} for r in data])
            provider = Mock()
            provider.generate.side_effect = translate
            engine.provider = lambda *args, **kwargs: provider
            result = engine.run_file(FileJob(source))
            self.assertEqual(len(requests), 1)
            self.assertEqual(len(requests[0]), 47)
            self.assertEqual({r['page'] for r in requests[0]}, set(range(1, 48)))
            self.assertEqual(result['status'], 'done')
            output = Presentation(result['output'])
            self.assertEqual([s.shapes[0].text for s in output.slides], [f'Current {i}' for i in range(47)])
            progress = [e['page'] for e in events if e['type'] == 'progress' and e['node'] == 'translate']
            self.assertEqual(progress, [0, 47])

    def test_batch_budget_counts_serialized_style_markers(self):
        records = [TextRecord(i, i + 1, '', '짧음',
                   '[[STYLE_0]]짧[[/STYLE_0]][[STYLE_1]]음[[/STYLE_1]]') for i in range(3)]
        size = len(json.dumps([{'id': r.id, 'page': r.page, 'text': r.styled_text} for r in records[:2]], ensure_ascii=False))
        batches = list(make_batches(records, size))
        self.assertEqual([len(b) for b in batches], [2, 1])
        self.assertEqual([r.id for b in batches for r in b], [0, 1, 2])
        self.assertEqual(len(list(make_batches(records, DEFAULT_BATCH_CHARS))), 1)
