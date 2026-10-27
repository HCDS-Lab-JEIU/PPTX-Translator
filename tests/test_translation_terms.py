"""English gloss reuse saves calls while retaining distinct parenthetical content."""
import json
import unittest
from unittest.mock import Mock

from pptx import Presentation
from pptx.util import Cm
from apply_translation import replace_frame
from pptx_pipeline.ai import AIResponse
from pptx_styles import styled_frame, plain_text
from translate_texts import PPTXTextTranslator, TextRecord
from translation_terms import reuse_english_label


class TranslationTermTests(unittest.TestCase):
    def test_complete_labels_reuse_existing_english_without_model_calls(self):
        for recover in (False, True):
            provider = Mock()
            records = [TextRecord(0,1,'','전류원 (Current Source)'),
                       TextRecord(1,2,'',"옴의 법칙(Ohm's Law)"),
                       TextRecord(2,2,'','직류（DC）')]
            translator = PPTXTextTranslator(provider, recover=recover)
            self.assertEqual(translator.translate_batch(records), {0:'Current Source',1:"Ohm's Law",2:'DC'})
            provider.generate.assert_not_called()
            self.assertEqual(translator.failures, [])

    def test_mixed_batch_sends_only_records_needing_translation(self):
        provider = Mock()
        provider.generate.return_value = AIResponse(json.dumps({'translations':[{'id':9,'text':'Measure current'}]}))
        translator = PPTXTextTranslator(provider, recover=True)
        output = translator.translate_batch([TextRecord(3,1,'','전압원 (Voltage Source)'),
                                            TextRecord(9,1,'','전류를 측정하세요')])
        self.assertEqual(output, {3:'Voltage Source',9:'Measure current'})
        payload = json.loads(provider.generate.call_args.args[0].split('Input records:\n')[1])
        self.assertEqual([r['id'] for r in payload], [9])

    def test_multiline_labels_keep_paragraph_boundaries(self):
        record = TextRecord(0,1,'','전류 (Current)\n\n전압 (Voltage)\n')
        self.assertEqual(reuse_english_label(record), 'Current\n\nVoltage\n')

    def test_reused_english_retains_its_own_emphasis(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(Cm(1),Cm(1),Cm(10),Cm(3)).text_frame
        run = frame.paragraphs[0].add_run()
        run.text = '전류원 ('
        run.font.bold = False
        english = frame.paragraphs[0].add_run()
        english.text = 'Current Source'
        english.font.bold = True
        frame.paragraphs[0].add_run().text = ')'
        styled, _ = styled_frame(frame)
        provider = Mock()
        result = PPTXTextTranslator(provider,recover=True).translate_batch([TextRecord(0,1,'',frame.text,styled)])
        self.assertEqual(plain_text(result[0]), 'Current Source')
        replace_frame(frame, result[0], recover=True)
        self.assertEqual(frame.text, 'Current Source')
        self.assertTrue(all(run.font.bold for run in frame.paragraphs[0].runs))
        provider.generate.assert_not_called()

    def test_partial_sentence_drops_only_literal_duplicate_gloss(self):
        source = TextRecord(0,1,'','전류원 (Current Source)을 사용하세요')
        for recover in (False, True):
            provider = Mock()
            provider.generate.return_value = AIResponse(json.dumps({'translations':[
                {'id':0,'text':'Use a Current Source (Current Source).'}]}))
            result = PPTXTextTranslator(provider,recover=recover).translate_batch([source])
            self.assertEqual(result[0], 'Use a Current Source.')
            self.assertEqual(provider.generate.call_count, 1)

    def test_duplicate_across_style_spans_preserves_surviving_style(self):
        record = TextRecord(0,1,'','전류원 (Current Source)을 사용하세요',
            '[[STYLE_0]]전류원 [[/STYLE_0]][[STYLE_1]](Current Source)을 사용하세요[[/STYLE_1]]')
        provider = Mock()
        provider.generate.return_value = AIResponse(json.dumps({'translations':[{'id':0,'text':
            '[[STYLE_0]]Use a Current Source[[/STYLE_0]][[STYLE_1]] (Current Source).[[/STYLE_1]]'}]}))
        translator = PPTXTextTranslator(provider,recover=True)
        result = translator.translate_batch([record])
        self.assertEqual(plain_text(result[0]), 'Use a Current Source.')
        self.assertIn('[[STYLE_0]]Use a Current Source[[/STYLE_0]]', result[0])
        self.assertEqual(translator.failures, [])

    def test_conditions_units_variables_numbers_and_equations_are_not_local_labels(self):
        for source in ('전압 (V)', '전류 (I)', '주파수 (kHz)', '측정 조건 (when heated)',
                       '제2 법칙 (Second Law)', '전압 ([[EQ_1]])', '전류 10 mA (Current)',
                       '전압 (Voltage)을 측정하세요'):
            with self.subTest(source=source):
                self.assertIsNone(reuse_english_label(TextRecord(0,1,'',source)))

    def test_nonenglish_target_never_uses_english_shortcut_or_rule(self):
        provider = Mock()
        provider.generate.return_value = AIResponse(json.dumps({'translations':[{'id':0,'text':'전류원'}]}))
        translator = PPTXTextTranslator(provider,target_language='Korean',recover=True)
        result = translator.translate_batch([TextRecord(0,1,'','전류원 (Current Source)')])
        self.assertEqual(result[0], '전류원')
        provider.generate.assert_called_once()
        self.assertNotIn('Never produce English (English)', provider.generate.call_args.args[0])
