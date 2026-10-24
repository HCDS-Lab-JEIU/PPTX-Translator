import base64
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Cm, Pt

from apply_translation import apply_translation, match_translations, replace_frame
import main
from translate_texts import read_records, write_translation


class ApplyTranslationTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.source = self.folder / "source.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        shape = group.shapes.add_textbox(Cm(2), Cm(3), Cm(8), Cm(3))
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.level = 1
        korean = paragraph.add_run()
        korean.text = "저항"
        korean.font.bold = True
        korean.font.size = Pt(24)
        korean.font.color.rgb = RGBColor(255, 0, 0)
        formula = paragraph.add_run()
        formula.text = " R = 10 Ω"
        formula.font.italic = True
        shape.text_frame.add_paragraph().text = "두 번째 줄"
        group.left = Cm(4)
        group.width = Cm(10)
        english = slide.shapes.add_textbox(Cm(2), Cm(7), Cm(8), Cm(1))
        english.text = "Keep this English"
        footer = slide.shapes.add_textbox(Cm(2), Cm(12), Cm(8), Cm(1))
        footer.text = "제외된 꼬리말"
        footer.element.nvSpPr.nvPr.get_or_add_ph().type = PP_PLACEHOLDER.FOOTER
        slide.notes_slide.notes_text_frame.text = "원본 발표자 노트"
        slide2 = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide2.shapes.add_table(1, 2, Cm(2), Cm(3), Cm(8), Cm(2)).table
        table.cell(0, 0).text = "전류"
        table.cell(0, 1).text = "2 A"
        blob = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII="
        )
        slide2.shapes.add_picture(BytesIO(blob), Cm(2), Cm(7), Cm(2), Cm(2))
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(self.source)
        with ZipFile(self.source, "a", compression=ZIP_DEFLATED) as archive:
            archive.writestr("custom/untouched.bin", b"arbitrary original package content")
        self.document = main.extract_pptx(self.source, self.folder / "output")
        self.manifest = self.folder / "content.json"
        self.save_manifest()
        self.texts = self.folder / "texts.txt"
        self.texts.write_text(main.format_text(self.document), encoding="utf-8")
        self.records = read_records(self.texts)
        self.translations = self.folder / "translated.txt"
        replacements = {"저항 R = 10 Ω\n두 번째 줄": "Resistance R = 10 Ω\nSecond line",
                        "전류": "Current", "Keep this English": "Model changed English"}
        write_translation(self.translations, self.records,
                          {r.id: (r.styled_text.replace("저항", "Resistance").replace("두 번째 줄", "Second line")
                                  if r.styled_text else replacements.get(r.text, r.text)) for r in self.records})
        self.destination = self.folder / "translated.pptx"

    def save_manifest(self):
        self.manifest.write_text(json.dumps(self.document, ensure_ascii=False), encoding="utf-8")

    def apply(self):
        return apply_translation(self.source, self.manifest, self.translations, self.destination)

    def test_copy_replaces_only_korean_and_keeps_other_package_parts_byte_identical(self):
        before = self.source.read_bytes()
        report = self.apply()
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual(report["replaced"], 2)
        self.assertEqual(report["kept_non_korean"], 2)
        self.assertEqual(report["modified_slides"], [1, 2])
        prs = Presentation(self.destination)
        self.assertEqual(len(prs.slides), 3)
        self.assertEqual(prs.slides[0].shapes[0].shapes[0].text, "Resistance R = 10 Ω\nSecond line")
        self.assertEqual(prs.slides[0].shapes[1].text, "Keep this English")
        self.assertEqual(prs.slides[0].shapes[2].text, "제외된 꼬리말")
        self.assertEqual(prs.slides[1].shapes[0].table.cell(0, 0).text, "Current")
        self.assertEqual(prs.slides[1].shapes[0].table.cell(0, 1).text, "2 A")
        modified = {"ppt/slides/slide1.xml", "ppt/slides/slide2.xml"}
        with ZipFile(self.source) as original, ZipFile(self.destination) as translated:
            self.assertEqual(original.namelist(), translated.namelist())
            for name in original.namelist():
                if name not in modified:
                    self.assertEqual(original.read(name), translated.read(name), name)
            for name in modified:
                old, new = etree.fromstring(original.read(name)), etree.fromstring(translated.read(name))
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for root in (old, new):
                    for node in root.xpath("//a:t", namespaces=ns):
                        node.text = ""
                        node.attrib.clear()
                self.assertEqual(etree.tostring(old), etree.tostring(new))

    def test_old_translation_without_shape_ids_still_matches(self):
        text = self.translations.read_text(encoding="utf-8")
        import re
        self.translations.write_text(re.sub(r"; shape=\d+", "", text), encoding="utf-8")
        self.assertEqual(self.apply()["replaced"], 2)

    def test_wrong_source_hash_does_not_replace_existing_destination(self):
        self.document["source_sha256"] = "0" * 64
        self.save_manifest()
        self.destination.write_bytes(b"previous output")
        with self.assertRaisesRegex(ValueError, "추출 당시"):
            self.apply()
        self.assertEqual(self.destination.read_bytes(), b"previous output")

    def test_stale_legacy_manifest_text_is_rejected(self):
        self.document.pop("source_sha256")
        self.document["slides"][0]["texts"][0]["text"] = "outdated source text"
        self.save_manifest()
        with self.assertRaisesRegex(ValueError, "원문 텍스트"):
            self.apply()
        self.assertFalse(self.destination.exists())

    def test_duplicate_and_missing_translation_rejected(self):
        text = self.translations.read_text(encoding="utf-8")
        for invalid in (text + text.splitlines()[0] + "\n", "\n".join(text.splitlines()[1:])):
            self.translations.write_text(invalid, encoding="utf-8")
            with self.assertRaises(ValueError):
                self.apply()
        self.assertFalse(self.destination.exists())

    def test_ambiguous_legacy_positions_rejected(self):
        document = deepcopy(self.document)
        record = deepcopy(document["slides"][0]["texts"][0])
        record["shape_id"] += 100
        document["slides"][0]["texts"].append(record)
        import re
        self.translations.write_text(re.sub(r"; shape=\d+", "", self.translations.read_text(encoding="utf-8")), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "위치가 중복"):
            match_translations(document, read_records(self.translations))

    def test_soft_break_and_empty_soft_line(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(0, 0, Cm(5), Cm(5)).text_frame
        frame.text = "첫 줄\v\v마지막"
        self.assertFalse(replace_frame(frame, "First\nMiddle\nLast"))
        self.assertEqual(frame.text, "First\vMiddle\vLast")

    def test_different_paragraph_count_uses_original_paragraph_style(self):
        prs = Presentation(self.source)
        frame = prs.slides[0].shapes[0].shapes[0].text_frame
        self.assertTrue(replace_frame(frame, "One\nTwo\nThree"))
        self.assertEqual(frame.text, "One\nTwo\nThree")
        self.assertEqual(frame.paragraphs[0].runs[0].font.size, Pt(24))
        self.assertEqual(frame.paragraphs[0].level, 1)


if __name__ == "__main__":
    unittest.main()
