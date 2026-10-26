import base64
from copy import deepcopy
from io import BytesIO, StringIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from contextlib import redirect_stdout, redirect_stderr
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Cm, Pt

import main
from legacy import pipeline
from pptx_pipeline.ai import AIResponse, ProviderRegistry, image_read_evidence
from pptx_pipeline.layout import (LayoutPolicy, PPTXLayoutModifier, inspect_layout,
    validate_plan, apply_layout_changes, text_and_math_signature, LayoutPlanner)
from pptx_pipeline.rendering import RenderResult
from pptx_pipeline.stages import (AnalysisStage, TranslationStage, ApplicationStage,
    ModificationStage, PresentationPipeline, PipelineContext)
from tests.test_native_math import native_equation
from gemini_cli import EmptyModelResponseError, GeminiCLIError

PIXEL = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII=')


class FakeRenderer:
    def __init__(self):
        self.calls = []
        self.requested_pages = []

    def render(self, source, directory, *, pages=None):
        self.calls.append(Path(source))
        self.requested_pages.append(pages)
        directory = Path(directory)
        directory.mkdir(parents=True, exist_ok=True)
        prs = Presentation(source)
        slides, images = {}, {}
        for page, slide in enumerate(prs.slides, 1):
            if pages and page not in pages:
                continue
            image = directory / f'slide_{page:03d}.png'
            image.write_bytes(PIXEL)
            images[page] = image
            metrics = []
            for shape, _ in main.walk_shapes(slide.shapes):
                if shape.has_text_frame:
                    font = shape.text_frame.paragraphs[0].font.size
                    size = font.pt if font is not None else 20
                    metrics.append({'shape_id': shape.shape_id, 'row': None, 'column': None,
                                    'base_font_pt': size, 'overflow': size > 17})
            slides[page] = {'page': page, 'text_metrics': metrics}
        return RenderResult(images, slides)


class FakeProvider:
    model = 'test-model'

    def __init__(self):
        self.visual_calls = 0
        self.text_calls = 0

    def generate(self, prompt, *, model=None, images=()):
        if not images:
            self.text_calls += 1
            data = json.loads(prompt.split('Input records:\n', 1)[1])
            return AIResponse(json.dumps({'translations': [
                {'id': item['id'], 'text': item['text'].replace('설명', 'Description')} for item in data]}))
        self.visual_calls += 1
        first = self.visual_calls == 1
        return AIResponse(json.dumps({'reviews': [{'page': 1, 'needs_changes': first,
            'issues': ['Overflow'] if first else []}], 'changes': [
                {'page': 1, 'shape_id': 2, 'font_scale': 0.8, 'reason': 'Fit translation'}] if first else []}),
                tuple(map(Path, images)))


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.source = self.folder / 'source.pptx'
        prs = Presentation()
        prs.slide_width, prs.slide_height = Cm(30), Cm(20)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Cm(2), Cm(3), Cm(10), Cm(3))
        shape.text = '설명 '
        shape.text_frame.paragraphs[0].font.size = Pt(20)
        shape.text_frame.paragraphs[0]._p.append(native_equation())
        slide.shapes.add_picture(BytesIO(PIXEL), Cm(15), Cm(8), Cm(4), Cm(2))
        prs.save(self.source)
        with ZipFile(self.source, 'a') as archive:
            archive.writestr('custom/keep.bin', b'original custom part')
        self.document = main.extract_pptx(self.source, self.folder / 'output')
        self.manifest = self.folder / 'content.json'
        self.manifest.write_text(json.dumps(self.document), encoding='utf-8')
        self.renderer = FakeRenderer()
        self.render = self.renderer.render(self.source, self.folder / 'before')
        self.policy = LayoutPolicy(max_passes=1)
        self.snapshot = inspect_layout(self.source, self.render, self.document, self.policy)

    def plan(self, change=None):
        return {'reviews': [{'page': 1, 'needs_changes': True, 'issues': ['overflow']}],
                'changes': [change or {'page': 1, 'shape_id': 2, 'font_scale': 0.8, 'reason': 'fit'}]}

    def test_entire_pipeline_with_replaceable_ai_and_renderer(self):
        original = self.source.read_bytes()
        provider = FakeProvider()
        modifier = PPTXLayoutModifier(provider, self.renderer, policy=self.policy, overview=False)
        flow = PresentationPipeline([AnalysisStage(), TranslationStage(provider), ApplicationStage(), ModificationStage(modifier)])
        context = PipelineContext(self.source, self.folder / 'flow')
        result = flow.run(context)
        self.assertEqual(result['status'], 'completed')
        self.assertEqual(list(result['stages']), ['analyze', 'translate', 'apply', 'modify'])
        self.assertEqual(provider.text_calls, 1)
        self.assertEqual(provider.visual_calls, 2)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertTrue(context.modified_pptx.is_file())
        self.assertIn('Description', Presentation(context.modified_pptx).slides[0].shapes[0].text)
        self.assertEqual(Presentation(context.modified_pptx).slides[0].shapes[0].text_frame.paragraphs[0].font.size, Pt(16))

    def test_unchanged_pages_do_not_repeat_visual_review(self):
        prs = Presentation(self.source)
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(self.source)
        document = main.extract_pptx(self.source, self.folder / 'two-pages')
        self.manifest.write_text(json.dumps(document), encoding='utf-8')
        class MultiPageProvider:
            def __init__(self): self.image_counts = []
            def generate(self, prompt, *, images=()):
                self.image_counts.append(len(images))
                first = len(self.image_counts) == 1
                pages = [1, 2] if first else [1]
                return AIResponse(json.dumps({'reviews': [{'page': p, 'needs_changes': first and p == 1,
                    'issues': []} for p in pages], 'changes': [{'page': 1, 'shape_id': 2,
                    'font_scale': 0.8, 'reason': 'fit'}] if first else []}), tuple(images))
        provider = MultiPageProvider()
        renderer = FakeRenderer()
        modifier = PPTXLayoutModifier(provider, renderer, policy=self.policy, pages_per_request=2, overview=False)
        report = modifier.modify(self.source, self.folder / 'selective.pptx', reference=self.source,
                                 manifest=self.manifest, workdir=self.folder / 'selective-work')
        self.assertEqual(provider.image_counts, [4, 2])
        self.assertEqual(renderer.requested_pages[-1], [1])
        self.assertEqual(report['status'], 'verified')
        self.assertEqual(report['reviewed_pages'], [1, 2])

    def test_empty_review_skips_page_and_saves_each_response(self):
        prs = Presentation(self.source)
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(self.source)
        document = main.extract_pptx(self.source, self.folder / 'skip-output')
        self.manifest.write_text(json.dumps(document), encoding='utf-8')
        work = self.folder / 'skip-work'
        class EmptyFirstProvider:
            calls = 0
            def generate(inner, prompt, *, images=()):
                inner.calls += 1
                if inner.calls == 1:
                    raise EmptyModelResponseError('empty')
                self.assertTrue((work / 'review_01_001.json').is_file())
                self.assertTrue((work / 'progress.json').is_file())
                self.assertEqual(len(Presentation(work / 'preview.pptx').slides), 2)
                return AIResponse(json.dumps({'reviews': [{'page': 2, 'needs_changes': False,
                                                          'issues': []}], 'changes': []}), tuple(images))
        provider = EmptyFirstProvider()
        report = PPTXLayoutModifier(provider, FakeRenderer(), overview=False, pages_per_request=1).modify(self.source,
            self.folder / 'skip.pptx', reference=self.source, manifest=self.manifest, workdir=work)
        self.assertEqual(provider.calls, 2)
        self.assertEqual(report['status'], 'needs_review')
        self.assertEqual(report['reviewed_pages'], [2])
        self.assertEqual(report['unreviewed_pages'], [1])
        self.assertEqual(report['failures'][0]['pages'], [1])
        self.assertTrue((self.folder / 'skip.pptx').exists())

    def test_license_errors_still_stop_review(self):
        class DeniedProvider:
            def generate(self, *args, **kwargs):
                raise GeminiCLIError('PERMISSION_DENIED')
        with self.assertRaises(GeminiCLIError):
            PPTXLayoutModifier(DeniedProvider(), FakeRenderer(), overview=False).modify(self.source,
                self.folder / 'denied.pptx', reference=self.source, manifest=self.manifest,
                workdir=self.folder / 'denied-work')

    def test_cli_uses_medium_translation_and_codex_review_by_default(self):
        with patch('legacy.pipeline.ProviderRegistry.create', return_value=FakeProvider()) as create:
            with patch('legacy.pipeline.PresentationPipeline.run', return_value={'status': 'completed', 'stages': {}}):
                with redirect_stdout(StringIO()), redirect_stderr(StringIO()):
                    self.assertEqual(pipeline.main([str(self.source)]), 0)
        self.assertEqual([c.kwargs['model'] for c in create.call_args_list],
                         ['gemini-3.8-flash-medium', None])
        self.assertEqual([c.args[0] for c in create.call_args_list], ['gemini-cli', 'codex-app-server'])
        self.assertEqual(create.call_args_list[1].kwargs['effort'], 'low')

    def test_edit_keeps_formula_text_media_and_unknown_parts(self):
        output = self.folder / 'modified.pptx'
        plan = self.plan()
        plan['changes'].append({'page': 1, 'shape_id': 3, 'reason': 'Make room',
                               'box_cm': {'x': 16, 'y': 9, 'width': 3.2, 'height': 1.6}})
        apply_layout_changes(self.source, output, plan, self.snapshot, self.policy, self.snapshot)
        self.assertEqual(text_and_math_signature(Presentation(self.source)), text_and_math_signature(Presentation(output)))
        with ZipFile(self.source) as old, ZipFile(output) as new:
            for name in old.namelist():
                if name != 'ppt/slides/slide1.xml':
                    self.assertEqual(old.read(name), new.read(name))

    def test_whole_table_resize_updates_grid_dimensions(self):
        prs = Presentation(self.source)
        table_shape = prs.slides[0].shapes.add_table(2, 2, Cm(2), Cm(12), Cm(8), Cm(3))
        table_shape.table.cell(0, 0).text = '표'
        shape_id = table_shape.shape_id
        prs.save(self.source)
        document = main.extract_pptx(self.source, self.folder / 'table-output')
        rendered = self.renderer.render(self.source, self.folder / 'table-before')
        snapshot = inspect_layout(self.source, rendered, document, self.policy)
        plan = self.plan({'page': 1, 'shape_id': shape_id, 'box_cm': {
            'x': 2, 'y': 12, 'width': 9, 'height': 3.3}, 'reason': 'Fit table'})
        output = self.folder / 'resized-table.pptx'
        apply_layout_changes(self.source, output, plan, snapshot, self.policy, snapshot)
        table = Presentation(output).slides[0].shapes[-1].table
        self.assertEqual(sum(c.width for c in table.columns), Cm(9))
        self.assertEqual(sum(r.height for r in table.rows), round(3.3 * 360000))

    def test_rejects_bad_plans(self):
        invalid = [
            {'page': 1, 'shape_id': 2, 'font_scale': 0.1, 'reason': 'tiny'},
            {'page': 1, 'shape_id': 2, 'font_scale': float('nan'), 'reason': 'nan'},
            {'page': 1, 'shape_id': 200, 'font_scale': 0.8, 'reason': 'unknown'},
            {'page': 1, 'shape_id': 2, 'text': 'rewritten', 'reason': 'rewrite'},
            {'page': 1, 'shape_id': 3, 'box_cm': {'x': 15, 'y': 8, 'width': 2, 'height': 2}, 'reason': 'distort'},
            {'page': 1, 'shape_id': 3, 'box_cm': {'x': 29, 'y': 8, 'width': 4, 'height': 2}, 'reason': 'outside'},
        ]
        for change in invalid:
            with self.subTest(change=change), self.assertRaises(ValueError):
                validate_plan(self.plan(change), self.snapshot, self.policy, self.snapshot)

    def test_cumulative_shrink_and_readonly_geometry_are_rejected(self):
        current = deepcopy(self.snapshot)
        current[1]['shapes'][0]['frames'][0]['base_font_pt'] = 14
        with self.assertRaises(ValueError):
            validate_plan(self.plan(), current, self.policy, self.snapshot)
        locked = deepcopy(self.snapshot)
        locked[1]['shapes'][1]['geometry_editable'] = False
        with self.assertRaises(ValueError):
            validate_plan(self.plan({'page': 1, 'shape_id': 3, 'box_cm': {
                'x': 16, 'y': 8, 'width': 4, 'height': 2}, 'reason': 'move'}), locked, self.policy, self.snapshot)

    def test_no_image_read_no_edit(self):
        class BlindProvider:
            def generate(self, *args, **kwargs):
                return AIResponse('{}')
        with self.assertRaises(RuntimeError):
            LayoutPlanner(BlindProvider(), self.policy).propose(self.snapshot, self.snapshot,
                                                             self.render.images, self.render.images)

    def test_missing_review_and_duplicate_changes_rejected(self):
        plan = self.plan()
        plan['reviews'] = []
        with self.assertRaises(ValueError):
            validate_plan(plan, self.snapshot, self.policy, self.snapshot)
        plan = self.plan()
        plan['changes'] *= 2
        with self.assertRaises(ValueError):
            validate_plan(plan, self.snapshot, self.policy, self.snapshot)

    def test_registry_extensions_and_unknown_provider(self):
        registry = ProviderRegistry()
        registry.register('future-adapter', lambda **options: FakeProvider())
        self.assertIsInstance(registry.create('future-adapter'), FakeProvider)
        with self.assertRaises(ValueError):
            registry.create('not-implemented')

    def test_only_successful_image_reads_count_as_evidence(self):
        image = self.render.images[1]
        event = {'step_update': {'state': 'DONE', 'tool_name': 'view_file', 'tool_info': {
            'parameters': {'AbsolutePath': str(image)}, 'output': 'image shown'}}}
        self.assertEqual(image_read_evidence([event], [image], self.folder), (image.resolve(),))
        event['step_update']['tool_info']['error'] = 'denied'
        self.assertEqual(image_read_evidence([event], [image], self.folder), ())

    def test_failed_stage_records_error_and_stops(self):
        calls = []
        class Stage:
            def __init__(self, name, fail=False): self.name, self.fail = name, fail
            def run(self, context):
                calls.append(self.name)
                if self.fail: raise RuntimeError('failure')
                return {'ok': True}
        context = PipelineContext(self.source, self.folder / 'failed')
        with self.assertRaises(RuntimeError):
            PresentationPipeline([Stage('first'), Stage('second', True), Stage('third')]).run(context)
        state = json.loads((context.output_dir / 'pipeline.json').read_text(encoding='utf-8'))
        self.assertEqual(calls, ['first', 'second'])
        self.assertEqual(state['status'], 'failed')
        self.assertEqual(state['active_stage'], 'second')

    def test_analysis_cli_does_not_require_an_ai_backend(self):
        with patch('legacy.pipeline.ProviderRegistry.create', side_effect=AssertionError('AI should not load')):
            with redirect_stdout(StringIO()), redirect_stderr(StringIO()):
                result = pipeline.main([str(self.source), '-o', str(self.folder / 'analysis'), '--stop-after', 'analyze'])
        self.assertEqual(result, 0)


if __name__ == '__main__':
    unittest.main()
