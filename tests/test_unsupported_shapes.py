"""Unsupported contentPart nodes must survive while editable slides keep working."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.util import Cm

from apply_translation import apply_translation
from main import extract_pptx, format_text, walk_shapes
from pptx_pipeline.layout import LayoutPolicy, apply_layout_changes, inspect_layout
from pptx_pipeline.ai import AIResponse
from pptx_gui.engine import DesktopEngine, RunControl
from pptx_gui.models import AppSettings, FileJob
from pptx_text import MC, P
from tests.test_pipeline import FakeRenderer, PIXEL
from translate_texts import read_records, write_translation


def content_part():
    return etree.fromstring(('''<p:contentPart xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdInk">
        <p:extLst><p:ext uri="preserve-unsupported-content"/></p:extLst></p:contentPart>''').encode())


class UnsupportedShapeTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.pptx'
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.shapes.add_textbox(Cm(2), Cm(3), Cm(10), Cm(3)).text = '전류'

    def save(self):
        self.prs.save(self.source)
        with ZipFile(self.source, 'a') as archive:
            archive.writestr('ppt/ink/preserved.bin', b'original embedded content')

    def test_direct_content_part_does_not_abort_extraction(self):
        self.slide.shapes._spTree.append(content_part())
        self.save()
        original = self.source.read_bytes()
        document = extract_pptx(self.source, self.root / 'extract')
        self.assertEqual([r['text'] for r in document['slides'][0]['texts']], ['전류'])
        self.assertEqual(document['preserved_unsupported_shapes'], [{'page':1,'type':'contentPart'}])
        self.assertTrue(any('contentPart' in warning for warning in document['warnings']))
        self.assertEqual(self.source.read_bytes(), original)

    def test_unsupported_choice_uses_supported_fallback_once(self):
        picture = self.slide.shapes.add_picture(BytesIO(PIXEL), Cm(3), Cm(8), Cm(3), Cm(3))
        wrapper = etree.Element('{' + MC + '}AlternateContent', nsmap={'mc':MC})
        choice = etree.SubElement(wrapper, '{' + MC + '}Choice', Requires='p14')
        choice.append(content_part())
        fallback = etree.SubElement(wrapper, '{' + MC + '}Fallback')
        self.slide.shapes._spTree.replace(picture.element, wrapper)
        fallback.append(picture.element)
        self.save()
        document = extract_pptx(self.source, self.root / 'extract')
        self.assertEqual(len(document['slides'][0]['images']), 1)
        self.assertEqual(len(document['preserved_unsupported_shapes']), 1)
        self.assertEqual(len(list(walk_shapes(Presentation(self.source).slides[0].shapes))), 2)

    def test_group_and_choice_keep_editable_text_and_ignore_duplicate_fallback(self):
        group = self.slide.shapes.add_group_shape()
        group.shapes.add_textbox(Cm(2), Cm(8), Cm(6), Cm(3)).text = '전압'
        group.element.append(content_part())
        wrapper = etree.Element('{' + MC + '}AlternateContent', nsmap={'mc':MC})
        choice = etree.SubElement(wrapper, '{' + MC + '}Choice', Requires='p14')
        fallback = etree.SubElement(wrapper, '{' + MC + '}Fallback')
        fallback.append(deepcopy(group.element))
        self.slide.shapes._spTree.replace(group.element, wrapper)
        choice.append(group.element)
        self.save()
        document = extract_pptx(self.source, self.root / 'extract')
        self.assertEqual([r['text'] for r in document['slides'][0]['texts']], ['전류', '전압'])
        self.assertEqual(len(document['preserved_unsupported_shapes']), 1)

    def test_translation_and_layout_keep_unknown_xml_and_package_parts(self):
        self.slide.shapes._spTree.append(content_part())
        self.save()
        original = self.source.read_bytes()
        document = extract_pptx(self.source, self.root / 'extract')
        document['translate_all'] = True
        manifest = self.root / 'content.json'
        manifest.write_text(json.dumps(document), encoding='utf-8')
        texts = self.root / 'texts.txt'
        texts.write_text(format_text(document), encoding='utf-8')
        translated = self.root / 'translated.txt'
        write_translation(translated, read_records(texts), {0:'Current'})
        applied = self.root / 'translated.pptx'
        apply_translation(self.source, manifest, translated, applied, recover=True)
        renderer = FakeRenderer()
        policy = LayoutPolicy()
        snapshot = inspect_layout(applied, renderer.render(applied, self.root / 'render'), document, policy)
        plan = {'reviews':[{'page':1,'needs_changes':True,'issues':['fit']}],
                'changes':[{'page':1,'shape_id':2,'font_scale':.8,'reason':'fit'}]}
        output = self.root / 'reviewed.pptx'
        apply_layout_changes(applied, output, plan, snapshot, policy, snapshot)
        with ZipFile(self.source) as source_zip, ZipFile(output) as output_zip:
            self.assertEqual(source_zip.namelist(), output_zip.namelist())
            for name in source_zip.namelist():
                if name != 'ppt/slides/slide1.xml':
                    self.assertEqual(source_zip.read(name), output_zip.read(name), name)
            for archive in (source_zip, output_zip):
                xml = etree.fromstring(archive.read('ppt/slides/slide1.xml'), etree.XMLParser(remove_blank_text=True))
                node = xml.find('.//{' + P + '}contentPart')
                self.assertIsNotNone(node)
                if archive is source_zip:
                    preserved = etree.tostring(node, method='c14n')
                else:
                    self.assertEqual(etree.tostring(node, method='c14n'), preserved)
        self.assertEqual(next(walk_shapes(Presentation(output).slides[0].shapes))[0].text, 'Current')
        self.assertEqual(self.source.read_bytes(), original)

    def test_desktop_saves_result_and_reports_preserved_content(self):
        self.slide.shapes._spTree.append(content_part())
        self.save()
        settings = AppSettings()
        settings.save.directory = str(self.root / 'results')
        settings.adjust.enabled = False
        events = []
        engine = DesktopEngine(settings, {}, events.append, RunControl())
        provider = Mock()
        provider.generate.return_value = AIResponse('{"translations":[{"id":0,"text":"Current"}]}')
        engine.provider = lambda *args, **kwargs: provider
        result = engine.run_file(FileJob(self.source))
        self.assertEqual(result['status'], 'needs_review')
        self.assertEqual(result['preserved_unsupported_shapes'], [{'page':1,'type':'contentPart'}])
        self.assertTrue(any(e['type'] == 'warning' and 'contentPart' in e['message'] for e in events))
        self.assertEqual(next(walk_shapes(Presentation(result['output']).slides[0].shapes))[0].text, 'Current')
