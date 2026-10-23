"""Extraction checks: filtering must retain body numbers and table contents."""

import base64
from contextlib import redirect_stdout, redirect_stderr
from io import BytesIO, StringIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Cm

import main


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def textbox(self, slide, text, x, y, width=3, height=0.4):
        box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(width), Cm(height))
        box.text = text
        return box

    def save_extract(self, presentation):
        path = self.folder / "input.pptx"
        presentation.save(path)
        return main.extract_pptx(path, self.folder / "output")

    def test_noise_filter_preserves_body_title_and_table(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Cm(30), Cm(20)
        for number in range(1, 5):
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "같은 제목"
            slide.shapes.title.top, slide.shapes.title.height = Cm(0.2), Cm(0.5)
            self.textbox(slide, "반복 머리말", 1, 0.1)
            self.textbox(slide, "반복 꼬리말", 1, 19)
            self.textbox(slide, f"{number} / 4", 27, 19)
            self.textbox(slide, "42", 10, 10)
            table = slide.shapes.add_table(1, 1, Cm(10), Cm(19), Cm(3), Cm(0.5)).table
            table.cell(0, 0).text = "123"
            explicit = self.textbox(slide, "명시적 꼬리말", 1, 8)
            ph = explicit.element.nvSpPr.nvPr.get_or_add_ph()
            ph.type = PP_PLACEHOLDER.FOOTER
        doc = self.save_extract(prs)
        for slide in doc["slides"]:
            texts = [r["text"] for r in slide["texts"]]
            self.assertEqual(set(texts), {"같은 제목", "42", "123"})
        self.assertEqual(doc["removed_text_count"], 16)

    def test_group_transform_and_merged_table(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(Cm(1), Cm(2), Cm(2), Cm(1))
        box.text = "그룹 텍스트"
        group.left, group.top = Cm(5), Cm(6)
        group.width, group.height = Cm(4), Cm(3)
        table = slide.shapes.add_table(2, 2, Cm(2), Cm(12), Cm(8), Cm(2)).table
        table.cell(0, 0).merge(table.cell(1, 1))
        table.cell(0, 0).text = "병합 셀"
        doc = self.save_extract(prs)
        texts = doc["slides"][0]["texts"]
        group_record = next(r for r in texts if r["text"] == "그룹 텍스트")
        self.assertEqual(group_record["position_cm"],
                         {"x": 5, "y": 6, "width": 4, "height": 3})
        cells = [r for r in texts if r["kind"] == "table_cell"]
        self.assertEqual(len(cells), 1)
        self.assertEqual((cells[0]["row_span"], cells[0]["column_span"]), (2, 2))
        self.assertEqual(cells[0]["position_cm"]["width"], 8)

    def test_original_images_deduplicated_and_cli_outputs(self):
        # A tiny PNG fixture: extraction must preserve its original bytes.
        blob = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII="
        )
        prs = Presentation()
        for page in range(2):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self.textbox(slide, '본문 "따옴표"\n두 번째 줄', 2, 4)
            slide.shapes.add_picture(BytesIO(blob), Cm(2), Cm(6), Cm(3), Cm(2))
        source = self.folder / "input.pptx"
        prs.save(source)
        output = self.folder / "out"
        stdout, stderr = StringIO(), StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            self.assertEqual(main.main([str(source), "-o", str(output)]), 0)
        self.assertEqual(stdout.getvalue(), (output / "texts.txt").read_text(encoding="utf-8"))
        doc = json.loads((output / "content.json").read_text(encoding="utf-8"))
        images = [image for slide in doc["slides"] for image in slide["images"]]
        self.assertEqual(len(images), 2)
        self.assertEqual(images[0]["file"], images[1]["file"])
        self.assertEqual((output / images[0]["file"]).read_bytes(), blob)
        self.assertEqual(len(list((output / "resources").iterdir())), 1)
        self.assertIn('\\n', stdout.getvalue())

    def test_rotated_group_and_blank_page(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Cm(30), Cm(20)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Cm(1), Cm(2), Cm(4), Cm(2)).text = "회전"
        group.left, group.top, group.rotation = Cm(5), Cm(6), 90
        prs.slides.add_slide(prs.slide_layouts[6])
        doc = self.save_extract(prs)
        pos = doc["slides"][0]["texts"][0]["position_cm"]
        self.assertEqual(pos, {"x": 6, "y": 5, "width": 2, "height": 4})
        self.assertEqual(doc["slides"][1], {"page": 2, "texts": [], "images": []})

    def test_invalid_input_has_nonzero_exit(self):
        for name in ("missing.pptx", "legacy.ppt", "broken.pptx"):
            path = self.folder / name
            if name == "broken.pptx":
                path.write_bytes(b"not a pptx")
            with redirect_stdout(StringIO()), redirect_stderr(StringIO()):
                self.assertEqual(main.main([str(path), "-o", str(self.folder / 'out')]), 1)


if __name__ == "__main__":
    unittest.main()
