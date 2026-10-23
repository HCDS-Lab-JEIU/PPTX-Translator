"""Reconstruction works after the source PPTX has been removed."""

import base64
from contextlib import redirect_stdout, redirect_stderr
from io import BytesIO, StringIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation
from pptx.util import Cm

from legacy import copy_pptx
import main


class RebuildTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.blob = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aD1sAAAAASUVORK5CYII="
        )
        prs = Presentation()
        prs.slide_width, prs.slide_height = Cm(30), Cm(20)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text = slide.shapes.add_textbox(Cm(2), Cm(3), Cm(8), Cm(2))
        text.text = '본문 "따옴표"\n다음 줄'
        image = slide.shapes.add_picture(BytesIO(self.blob), Cm(5), Cm(6), Cm(4), Cm(2))
        image.crop_left, image.crop_right = 0.1, 0.2
        image.rotation = 90
        prs.slides.add_slide(prs.slide_layouts[6])
        source = self.folder / "source.pptx"
        prs.save(source)
        self.extracted = self.folder / "extracted"
        self.doc = main.extract_pptx(source, self.extracted)
        self.manifest = self.extracted / "content.json"
        self.manifest.write_text(json.dumps(self.doc, ensure_ascii=False), encoding="utf-8")
        source.unlink()

    def test_roundtrip_without_source_preserves_positions_images_and_blank_page(self):
        destination = self.folder / "restored.pptx"
        report = copy_pptx.rebuild_pptx(self.manifest, destination)
        self.assertEqual((report["pages"], report["texts"], report["images"]), (2, 1, 1))
        prs = Presentation(destination)
        self.assertEqual((prs.slide_width, prs.slide_height), (Cm(30), Cm(20)))
        self.assertEqual(len(prs.slides), 2)
        self.assertEqual(len(prs.slides[1].shapes), 0)
        text, image = prs.slides[0].shapes
        self.assertEqual(text.text, '본문 "따옴표"\n다음 줄')
        self.assertEqual((text.left, text.top, text.width, text.height), (Cm(2), Cm(3), Cm(8), Cm(2)))
        self.assertEqual(image.image.blob, self.blob)
        self.assertEqual((image.left, image.top, image.width, image.height), (Cm(5), Cm(6), Cm(4), Cm(2)))
        self.assertEqual(image.rotation, 90)
        self.assertAlmostEqual(image.crop_left, 0.1)
        self.assertAlmostEqual(image.crop_right, 0.2)

    def test_legacy_manifest_and_cli_report(self):
        for page in self.doc["slides"]:
            for record in page["texts"] + page["images"]:
                record.pop("z_order")
        self.manifest.write_text(json.dumps(self.doc), encoding="utf-8")
        with redirect_stdout(StringIO()), redirect_stderr(StringIO()):
            self.assertEqual(copy_pptx.main([str(self.manifest)]), 0)
        self.assertTrue((self.extracted / "copy.pptx").is_file())
        report = json.loads((self.extracted / "copy.report.json").read_text(encoding="utf-8"))
        self.assertTrue(any("겹침 순서" in warning for warning in report["warnings"]))

    def test_missing_image_is_an_error_and_does_not_save_incomplete_deck(self):
        image = self.doc["slides"][0]["images"][0]
        (self.extracted / image["file"]).unlink()
        with redirect_stdout(StringIO()), redirect_stderr(StringIO()):
            self.assertEqual(copy_pptx.main([str(self.manifest)]), 1)
        self.assertFalse((self.extracted / "copy.pptx").exists())

    def test_table_cells_are_editable_text(self):
        record = self.doc["slides"][0]["texts"][0]
        record.update(kind="table_cell", row=1, column=2, row_span=1, column_span=1)
        self.manifest.write_text(json.dumps(self.doc), encoding="utf-8")
        destination = self.folder / "table.pptx"
        report = copy_pptx.rebuild_pptx(self.manifest, destination)
        prs = Presentation(destination)
        self.assertTrue(prs.slides[0].shapes[0].name.endswith("_r1_c2"))
        self.assertEqual(prs.slides[0].shapes[0].text, record["text"])
        self.assertTrue(any("표 셀" in warning for warning in report["warnings"]))

    def test_ambiguous_rotation_fallback_and_resource_boundary(self):
        image = self.doc["slides"][0]["images"][0]
        image["rotation_degrees"] = 45
        warnings = []
        *_, rotation = copy_pptx.image_frame(image, warnings)
        self.assertEqual(rotation, 0)
        self.assertTrue(warnings)
        with self.assertRaises(ValueError):
            copy_pptx.resource_path(self.extracted, "../source.pptx")


if __name__ == "__main__":
    unittest.main()
