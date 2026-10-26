"""Ordered repair strategies and content-preserving whitespace/spacing operations."""
import re

from pptx.util import Pt

from apply_translation import clean_text, replace_frame, replace_paragraph, validate_paragraph_order
from pptx_text import EQUATION, equation_lines, frame_text, paragraph_text, equation_map


STAGES = ('font', 'spacing', 'images', 'compression')
STAGE_LABELS = {'font':'폰트 크기', 'spacing':'공백·간격', 'images':'그림 크기·위치', 'compression':'문장 간소화'}
STAGE_FIELDS = {'font': {'font_scale'}, 'spacing': {'spacing_scale', 'whitespace_text'},
                'images': {'box_cm'}, 'compression': {'text'}}
STAGE_INSTRUCTIONS = {
    'font': 'Try ONLY a minimal font_scale change within readability limits. Do not move objects or rewrite text.',
    'spacing': 'Try ONLY spacing_scale or whitespace_text. spacing_scale reduces paragraph gaps, line spacing '
        'and internal margins. whitespace_text may normalize repeated spaces/tabs and remove empty paragraphs. '
        'whitespace_text is PLAIN TEXT without STYLE markers. Keep every nonempty line and every word, '
        'ordinary word separator, number, formula and semantic style.',
    'images': 'Try ONLY box_cm on editable PICTURES, preserving aspect ratio and content. Do not move text boxes.',
    'compression': 'Try ONLY a shorter equivalent text expression. Preserve ALL facts, conditions, negation, '
        'technical terms, numbers, variables, units, equations and semantic style. Do not change font or geometry.',
}


def eligible_pages(stage, snapshots, pages):
    def editable(shape):
        if stage == 'images':
            return shape['kind'] == 'picture' and shape['geometry_editable']
        return any(frame['font_editable'] if stage == 'font' else frame['text_editable']
                   for frame in shape['frames'])
    return [page for page in pages if any(editable(shape) for shape in snapshots[page]['shapes'])]


def whitespace_form(text):
    return tuple(re.sub(r'[ \t]+', ' ', line).strip() for line in text.replace('\v', '\n').split('\n') if line.strip())


def validate_whitespace(frame, text):
    original = frame['text']
    if not frame.get('text_editable') or not isinstance(text, str) or not text.strip():
        raise ValueError('공백 조정 대상 텍스트가 올바르지 않습니다.')
    if re.search(r'(?m)^\s*(?:def |class |import |#include|public |return |for\s*\(|if\s*\()', original):
        raise ValueError('코드 프레임의 공백은 변경할 수 없습니다.')
    if whitespace_form(original) != whitespace_form(text):
        raise ValueError('공백 단계에서는 단어·단일 단어 구분 공백·내용·비어 있지 않은 문단을 바꿀 수 없습니다.')
    if text.count('\n') > original.count('\n'):
        raise ValueError('공백 단계에서 줄을 추가할 수 없습니다.')
    if (EQUATION.search(original) or frame.get('has_soft_breaks')) and equation_lines(original) != equation_lines(text):
        raise ValueError('수식·명시적 줄바꿈 위치는 변경할 수 없습니다.')


def spacing_profile(frame):
    result = {f'margin_{side}': getattr(frame, f'margin_{side}').pt for side in ('left','right','top','bottom')}
    for index, paragraph in enumerate(frame.paragraphs):
        value = paragraph.line_spacing
        result[f'line_{index}'] = value if isinstance(value, float) else value.pt if value is not None else 1.0
        for name in ('space_before', 'space_after'):
            value = getattr(paragraph, name)
            result[f'{name}_{index}'] = value.pt if value is not None else 0.0
    return result


def scale_spacing(frame, scale, base_font_pt):
    for side in ('left','right','top','bottom'):
        name = f'margin_{side}'
        setattr(frame, name, round(getattr(frame, name) * scale))
    for paragraph in frame.paragraphs:
        for name in ('space_before', 'space_after'):
            value = getattr(paragraph, name)
            if value is not None:
                setattr(paragraph, name, Pt(value.pt * scale))
        value = paragraph.line_spacing
        empty = not paragraph_text(paragraph._p, equation_map(frame)).strip()
        if value is None or isinstance(value, float):
            current = 1.0 if value is None else value
            floor = min(current, .65 if empty else .95)
            paragraph.line_spacing = max(floor, current * scale)
        else:
            fonts = [run.font.size.pt for run in paragraph.runs if run.font.size is not None]
            font = max(fonts or [base_font_pt or value.pt])
            floor = min(value.pt, font * .95)
            paragraph.line_spacing = Pt(max(floor, value.pt * scale))


def apply_whitespace(frame, text):
    """Preserve original nonempty paragraph/run formatting when removing blank paragraphs."""
    source = frame_text(frame)
    if len(text.split('\n')) == len(source.replace('\v', '\n').split('\n')):
        replace_frame(frame, text)
        return
    paragraphs = list(frame.paragraphs)
    tokens = equation_map(frame)
    if tokens or any('\v' in paragraph_text(p._p, tokens) for p in paragraphs):
        raise ValueError('수식·명시적 줄바꿈 프레임의 문단 수는 유지해야 합니다.')
    selected = paragraphs
    nonempty = iter(p for p in selected if paragraph_text(p._p, tokens).strip())
    blank = iter(p for p in selected if not paragraph_text(p._p, tokens).strip())
    replacements = []
    for line in text.split('\n'):
        paragraph = next(nonempty if line.strip() else blank)
        replace_paragraph(paragraph._p, line, tokens)
        replacements.append(paragraph._p)
    body = frame._txBody
    insertion = list(body).index(selected[0]._p)
    for paragraph in selected:
        body.remove(paragraph._p)
    for offset, paragraph in enumerate(replacements):
        body.insert(insertion + offset, paragraph)
    validate_paragraph_order(frame)
    if clean_text(frame_text(frame)) != clean_text(text):
        raise ValueError('공백 조정 후 텍스트 검증 실패')
