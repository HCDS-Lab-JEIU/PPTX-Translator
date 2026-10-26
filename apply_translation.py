"""원본 PPTX를 복사하고 translated.txt의 영어를 원래 텍스트 위치에 반영한다."""

import argparse
from collections import defaultdict
from copy import deepcopy
from difflib import SequenceMatcher
import json
import math
from pathlib import Path
import re
import shutil
import sys
from tempfile import TemporaryDirectory
from zipfile import ZipFile, BadZipFile

from lxml import etree
from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

from main import clean_text, file_sha256, text_records, walk_shapes
from translate_texts import TextRecord, changed_numbers, read_records
from translation_validation import TranslationValidationError, prepare_translation
from pptx_text import EDITABLE_SHAPE_TAGS, frame_text, equation_map, paragraph_text, math_text_nodes, MC
from pptx_styles import MARKER, SPAN, plain_text, relaxed_styles, styled_frame, validate_styles

KOREAN = re.compile(r"[\u1100-\u11ff\u3130-\u318f\uac00-\ud7a3]")
LOCATION = re.compile(
    r"^\[page=\d+\]\[x=([^,]+),y=([^,]+),w=([^,]+),h=([^ ]+) cm(.*)\]$"
)


def identity(record):
    return (record["shape_id"], record["kind"], record.get("row"), record.get("column"), record.get("math_text_index"))


def position_key(page, position, table=None, row=None, column=None, math_text=None):
    values = tuple(round(float(position[key]), 4) for key in ("x", "y", "width", "height"))
    if not all(math.isfinite(value) for value in values):
        raise ValueError("위치 정보가 유한 숫자가 아닙니다.")
    return page, values, table, row, column, math_text


def translation_location(record):
    match = LOCATION.fullmatch(record.prefix)
    if match is None:
        raise ValueError(f"번역 항목 {record.id}: 위치 형식을 읽을 수 없습니다.")
    position = dict(zip(("x", "y", "width", "height"), map(float, match.groups()[:4])))
    extra = {}
    for field in match[5].split(";"):
        for token in field.strip().split(","):
            if not token:
                continue
            key, value = token.split("=", 1)
            if key not in {"shape", "table", "row", "col", "math_text"} or key in extra:
                raise ValueError(f"알 수 없거나 중복된 위치 필드: {key}")
            extra[key] = int(value)
    key = position_key(record.page, position, extra.get("table"), extra.get("row"), extra.get("col"), extra.get("math_text"))
    return key, extra.get("shape")


def match_translations(document, translated):
    by_position = defaultdict(list)
    expected = []
    for page in document["slides"]:
        for record in page["texts"]:
            key = position_key(page["page"], record["position_cm"],
                record["shape_id"] if "row" in record else None,
                record.get("row"), record.get("column"), record.get("math_text_index"))
            by_position[key].append((page["page"], record))
            expected.append((page["page"], identity(record)))
    seen = set()
    matches = []
    for translation in translated:
        key, shape_id = translation_location(translation)
        candidates = by_position.get(key, [])
        if shape_id is not None:
            candidates = [item for item in candidates if item[1]["shape_id"] == shape_id]
        if len(candidates) != 1:
            raise ValueError(f"번역 항목 {translation.id}: 대응 원본이 없거나 위치가 중복됩니다. "
                             "동일한 원본을 다시 추출·번역하여 도형 ID를 포함하세요.")
        page, record = candidates[0]
        key = (page, identity(record))
        if key in seen:
            raise ValueError(f"페이지 {page}: 동일 도형에 대한 번역이 중복됩니다.")
        seen.add(key)
        matches.append((page, record, translation.text))
    if seen != set(expected) or len(expected) != len(seen):
        raise ValueError("추출 텍스트와 번역 텍스트의 항목 수 또는 도형 ID가 일치하지 않습니다.")
    return matches


def text_insertion_index(parent, nodes=(), anchor=None):
    """DrawingML requires endParaRPr after all text, fields, breaks and equations."""
    end = parent.find(qn('a:endParaRPr'))
    limit = list(parent).index(end) if end is not None else len(parent)
    if nodes:
        return min(list(parent).index(nodes[0].getparent()), limit)
    if anchor is not None:
        return min(list(parent).index(anchor), limit)
    return limit


def validate_paragraph_order(frame):
    for paragraph in frame.paragraphs:
        children = list(paragraph._p)
        for index, child in enumerate(children):
            if child.tag == qn('a:endParaRPr') and index != len(children) - 1:
                raise ValueError('문단 끝 서식 뒤에 텍스트가 배치되어 표시할 수 없습니다.')


def replace_segment(nodes, translated, parent, anchor=None, styles=None, fallback=None):
    """Unchanged fragments retain their original run formatting."""
    if fallback is not None and not MARKER.search(translated):
        insertion = text_insertion_index(parent, nodes, anchor)
        template = max(nodes, key=lambda node: len(node.text or '')).getparent() if nodes else fallback
        run = deepcopy(template)
        run.tag = qn('a:r')
        run.attrib.clear()
        for child in list(run):
            if child.tag != qn('a:rPr'):
                run.remove(child)
        for owner in dict.fromkeys(node.getparent() for node in nodes):
            parent.remove(owner)
        text = OxmlElement('a:t')
        text.text = translated
        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        run.append(text)
        parent.insert(insertion, run)
        return
    if styles is not None and MARKER.search(translated):
        insertion = text_insertion_index(parent, nodes, anchor)
        for run in dict.fromkeys(node.getparent() for node in nodes):
            parent.remove(run)
        matches = list(SPAN.finditer(translated))
        for index, match in enumerate(matches):
            run = deepcopy(styles[match[1]])
            run.tag = qn("a:r")
            run.attrib.clear()
            for text in run.findall(qn("a:t")):
                run.remove(text)
            text = OxmlElement("a:t")
            start = 0 if index == 0 else match.start()
            end = matches[index + 1].start() if index + 1 < len(matches) else len(translated)
            text.text = plain_text(translated[start:end])
            text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            run.append(text)
            parent.insert(insertion + index, run)
        return
    if not nodes:
        if translated:
            run = OxmlElement("a:r")
            text = OxmlElement("a:t")
            text.text = translated
            run.append(text)
            end = anchor if anchor is not None else parent.find(qn("a:endParaRPr"))
            parent.insert(list(parent).index(end) if end is not None else len(parent), run)
        return
    chunks = [node.text or "" for node in nodes]
    original = "".join(chunks)
    owners = [index for index, chunk in enumerate(chunks) for _ in chunk]
    output = [""] * len(nodes)
    for operation, i1, i2, j1, j2 in SequenceMatcher(None, original, translated, autojunk=False).get_opcodes():
        if operation == "equal":
            for offset in range(i2 - i1):
                output[owners[i1 + offset]] += translated[j1 + offset]
        elif operation in {"replace", "insert"}:
            owner = owners[min(i1, len(owners) - 1)] if owners else 0
            output[owner] += translated[j1:j2]
    for node, value in zip(nodes, output):
        if node.text != value:
            node.text = value
            if value.startswith(" ") or value.endswith(" "):
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            if node.getparent().tag == qn("a:fld"):
                # A translated dynamic field must not revert on the next refresh.
                field = node.getparent()
                field.tag = qn("a:r")
                field.attrib.clear()


def replace_paragraph(paragraph, translated, tokens=None, styles=None, fallback=None):
    tokens = tokens or {}
    groups = [[]]
    breaks = []
    boundaries = []
    for child in paragraph:
        if child.tag == qn("a:br"):
            breaks.append(child)
            boundaries.append("\v")
            groups.append([])
        elif child in tokens:
            breaks.append(child)
            boundaries.append(tokens[child])
            groups.append([])
        elif child.tag in {qn("a:r"), qn("a:fld")}:
            groups[-1].extend(child.findall(qn("a:t")))
    parts = re.split(r"(\[\[EQ_\d+\]\]|\v)", translated)
    lines = parts[::2]
    if parts[1::2] != boundaries:
        raise ValueError("번역문의 수식 위치 또는 줄바꿈이 원본과 다릅니다.")
    if len(groups) != len(lines):
        raise ValueError("문단 내 줄바꿈 매핑이 올바르지 않습니다.")
    for index, (nodes, text) in enumerate(zip(groups, lines)):
        replace_segment(nodes, text, paragraph, breaks[index] if index < len(breaks) else None, styles, fallback)


def replace_frame(frame, translated, *, recover=False):
    styles = None
    fallback = None
    source, templates = styled_frame(frame)
    if MARKER.search(translated):
        if source is None:
            raise ValueError("원본에 없는 서식 표시입니다.")
        if not recover or not relaxed_styles(source, translated):
            validate_styles(source.replace("\v", "\n").strip("\n"), translated.replace("\v", "\n"))
        styles = {key: deepcopy(value) for key, value in templates.items()}
    elif recover and source:
        # No semantic span mapping is available. Use a consistent source base style,
        # never assign emphasis by translated character position or reject the prose.
        fallback = max(templates.values(), key=lambda run: len(''.join(run.itertext())))
    paragraphs = list(frame.paragraphs)
    tokens = equation_map(frame)
    active = [index for index, paragraph in enumerate(paragraphs) if paragraph_text(paragraph._p, tokens).strip()]
    if not active:
        raise ValueError("원본 텍스트 프레임이 비어 있습니다.")
    first, last = active[0], active[-1]
    selected = paragraphs[first:last + 1]
    counts = [len(paragraph_text(paragraph._p, tokens).split("\v")) for paragraph in selected]
    lines = translated.replace("\r\n", "\n").replace("\v", "\n").split("\n")
    changed_paragraphs = len(lines) != sum(counts)
    if tokens and changed_paragraphs:
        raise ValueError("수식 포함 텍스트의 문단 수가 달라 수식을 안전하게 유지할 수 없습니다. 다시 번역하세요.")
    if not changed_paragraphs:
        offset = 0
        for paragraph, count in zip(selected, counts):
            replace_paragraph(paragraph._p, "\v".join(lines[offset:offset + count]), tokens, styles, fallback)
            offset += count
    else:
        # A changed paragraph count has no exact style correspondence.
        body = frame._txBody
        insertion = list(body).index(selected[0]._p)
        replacements = []
        for index, line in enumerate(lines):
            template = deepcopy(selected[min(index, len(selected) - 1)]._p)
            runs = template.xpath('./a:r | ./a:fld')
            default_run = deepcopy(runs[0]) if runs else fallback
            for child in list(template):
                if child.tag in {qn("a:r"), qn("a:fld"), qn("a:br")}:
                    template.remove(child)
            replace_paragraph(template, line, styles=styles, fallback=default_run)
            replacements.append(template)
        for paragraph in selected:
            body.remove(paragraph._p)
        for offset, paragraph in enumerate(replacements):
            body.insert(insertion + offset, paragraph)
    validate_paragraph_order(frame)
    if clean_text(frame_text(frame)) != clean_text(plain_text(translated)):
        raise ValueError("텍스트 교체 후 내용 검증에 실패했습니다.")
    return changed_paragraphs


def sync_text_fallback(shape, expected, translated, *, recover=False):
    """Update equivalent editable fallback text; never alter bitmap fallbacks."""
    choice = next((node for node in shape.element.iterancestors()
                   if node.tag == "{" + MC + "}Choice"), None)
    if choice is None:
        return 0
    count = 0
    for fallback in choice.getparent().findall("{" + MC + "}Fallback"):
        ids = [node for node in fallback.iter(qn("p:cNvPr")) if node.get("id") == str(shape.shape_id)]
        for identifier in ids:
            element = identifier.getparent().getparent()
            if element.tag not in EDITABLE_SHAPE_TAGS:
                continue
            alternate = shape._parent._shape_factory(element)
            if "row" in expected:
                if not alternate.has_table:
                    continue
                row, column = expected["row"] - 1, expected["column"] - 1
                if row >= len(alternate.table.rows) or column >= len(alternate.table.columns):
                    continue
                frame = alternate.table.cell(row, column).text_frame
            elif alternate.has_text_frame:
                frame = alternate.text_frame
            else:
                continue
            if expected["kind"] == "math_text":
                nodes = math_text_nodes(frame)
                index = expected["math_text_index"] - 1
                if index < len(nodes) and nodes[index].text == expected["text"]:
                    nodes[index].text = translated
                    count += 1
            elif clean_text(frame_text(frame)) == expected["text"]:
                replace_frame(frame, translated, recover=True) if recover else replace_frame(frame, translated)
                count += 1
    return count


def apply_translation(source, manifest="output/content.json", translations="output/translated.txt", destination=None,
                      *, recover=False):
    source, manifest, translations = Path(source), Path(manifest), Path(translations)
    destination = Path(destination) if destination else translations.parent / "translated.pptx"
    if source.suffix.lower() != ".pptx" or destination.suffix.lower() != ".pptx":
        raise ValueError("입력과 출력은 .pptx 파일이어야 합니다.")
    if destination.resolve() in {source.resolve(), manifest.resolve(), translations.resolve()}:
        raise ValueError("출력 경로는 원본·입력 파일과 달라야 합니다.")
    document = json.loads(manifest.read_text(encoding="utf-8-sig"))
    if document.get("coordinate_unit") != "cm" or document.get("coordinate_origin") != "top-left":
        raise ValueError("cm 단위, 왼쪽 위 원점의 추출 데이터가 필요합니다.")
    matches = match_translations(document, read_records(translations))
    source_hash = file_sha256(source)
    if document.get("source_sha256") and document["source_sha256"] != source_hash:
        raise ValueError("원본 PPTX가 추출 당시 파일과 다릅니다.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    report = {"source": source.name, "source_sha256": source_hash, "matched": len(matches),
              "replaced": 0, "kept_non_korean": 0, "modified_slides": [], "paragraph_count_changed": [],
              "fallback_texts_updated": 0, "kept_original": [], "formatting_deferred": []}
    with TemporaryDirectory(prefix="pptx_translation_", dir=destination.parent) as temporary:
        copied = Path(temporary) / "original.pptx"
        staged = Path(temporary) / "translated.pptx"
        shutil.copy2(source, copied)
        prs = Presentation(str(copied))
        if len(prs.slides) != len(document["slides"]):
            raise ValueError("원본과 추출 데이터의 슬라이드 수가 다릅니다.")
        for dimension, actual in (("width", prs.slide_width), ("height", prs.slide_height)):
            if abs(actual / 360000 - document["slide_size_cm"][dimension]) > 0.0002:
                raise ValueError("원본 슬라이드 크기가 추출 데이터와 다릅니다.")
        actual_records, shapes = {}, {}
        for page, slide in enumerate(prs.slides, 1):
            for shape, matrix in walk_shapes(slide.shapes):
                shapes[(page, shape.shape_id)] = shape
                for record in text_records(shape, matrix, page):
                    actual_records[(page, identity(record))] = record
        if document.get("text_schema_version", 1) < 2 and any(record.get("equations") for record in actual_records.values()):
            raise ValueError("수식 지원 전 추출 데이터입니다. 원본을 다시 추출·번역하세요.")
        changed = set()
        for page, expected, translated in matches:
            actual = actual_records.get((page, identity(expected)))
            if actual is None or actual["text"] != expected["text"]:
                raise ValueError(f"페이지 {page}, 도형 {expected['shape_id']}: 원문 텍스트가 다릅니다.")
            if any(abs(actual["position_cm"][key] - expected["position_cm"][key]) > 0.0002
                   for key in ("x", "y", "width", "height")):
                raise ValueError(f"페이지 {page}, 도형 {expected['shape_id']}: 위치가 다릅니다.")
            if not document.get("translate_all") and not KOREAN.search(expected["text"]):
                report["kept_non_korean"] += 1
                continue
            shape = shapes[(page, expected["shape_id"])]
            if translated == expected['text']:
                # Preserve the XML verbatim for skipped/failed translations.
                continue
            notes = []
            if recover:
                record = TextRecord(0, page, '', expected['text'], actual.get('styled_text'))
                try:
                    translated, notes = prepare_translation(record, translated)
                    if changed_numbers(record, translated):
                        raise TranslationValidationError('numeric_mismatch')
                except TranslationValidationError as exc:
                    report['kept_original'].append({'page': page, 'shape_id': shape.shape_id, 'reason': exc.code})
                    continue
            if (actual.get("styled_text") and translated != expected["text"]
                    and not MARKER.search(translated)):
                if not recover:
                    raise ValueError("혼합 서식 번역에 의미별 서식 표시가 없습니다. 원본을 다시 추출·번역하세요.")
            frame = (shape.table.cell(expected["row"] - 1, expected["column"] - 1).text_frame
                     if "row" in expected else shape.text_frame)
            # Replacement can fail after mutating runs. Roll back the entire text body
            # and editable fallback before continuing to the next record.
            roots = [frame._txBody]
            choice = next((n for n in shape.element.iterancestors()
                           if n.tag == '{' + MC + '}Choice'), None)
            if choice is not None:
                roots.extend(choice.getparent().findall('{' + MC + '}Fallback'))
            backups = [(root, deepcopy(root)) for root in roots] if recover else []
            try:
                paragraph_changed = False
                if expected["kind"] == "math_text":
                    node = math_text_nodes(frame)[expected["math_text_index"] - 1]
                    node.text = translated
                else:
                    paragraph_changed = replace_frame(frame, translated, recover=True) if recover else replace_frame(frame, translated)
                fallback_count = sync_text_fallback(shape, expected, translated, recover=recover)
            except ValueError:
                if not recover:
                    raise
                for root, backup in backups:
                    root.attrib.clear()
                    root.attrib.update(backup.attrib)
                    root.text, root.tail = backup.text, backup.tail
                    root[:] = list(backup)
                report['kept_original'].append({'page': page, 'shape_id': shape.shape_id,
                    'reason': '서식·수식·줄바꿈 적용 검증 실패'})
                continue
            if paragraph_changed:
                report["paragraph_count_changed"].append({"page": page, "shape_id": shape.shape_id})
            if notes or paragraph_changed:
                report['formatting_deferred'].append({'page': page, 'shape_id': shape.shape_id,
                    'row': expected.get('row'), 'column': expected.get('column'),
                    'reasons': sorted(set(notes + (['paragraph_reflow_deferred'] if paragraph_changed else [])))})
            report["fallback_texts_updated"] += fallback_count
            changed.add(page)
            report["replaced"] += 1
        patches = {str(prs.slides[page - 1].part.partname).lstrip("/"):
                   etree.tostring(prs.slides[page - 1].element, encoding="UTF-8", xml_declaration=True, standalone=True)
                   for page in changed}
        # All unrelated package members retain their original bytes and ZIP metadata.
        if patches:
            with ZipFile(copied) as original, ZipFile(staged, "w") as updated:
                updated.comment = original.comment
                for entry in original.infolist():
                    updated.writestr(entry, patches.get(entry.filename, original.read(entry)))
        else:
            shutil.copy2(copied, staged)
        staged.replace(destination)
    report["modified_slides"] = sorted(changed)
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="원본 PPTX 경로")
    parser.add_argument("--manifest", type=Path, default=Path("output/content.json"))
    parser.add_argument("--translations", type=Path, default=Path("output/translated.txt"))
    parser.add_argument("-o", "--output", type=Path, help="기본: 번역 텍스트 폴더/translated.pptx")
    args = parser.parse_args(argv)
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    destination = args.output if args.output else args.translations.parent / "translated.pptx"
    try:
        report = apply_translation(args.source, args.manifest, args.translations, destination)
        destination.with_suffix(".report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except (OSError, ValueError, KeyError, TypeError, IndexError, BadZipFile, PackageNotFoundError) as exc:
        print(f"번역 반영 실패: {exc}", file=sys.stderr)
        return 1
    print(f"저장 완료: {destination.resolve()} (교체 {report['replaced']}개, "
          f"한글 없는 항목 유지 {report['kept_non_korean']}개)")
    if report["paragraph_count_changed"]:
        print("참고: 번역문 줄 수가 다른 일부 항목은 문단을 재구성했습니다. 보고서를 확인하세요.", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
