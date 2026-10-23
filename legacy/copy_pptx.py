"""content.json과 resources만 사용하여 편집 가능한 PPTX를 재구성한다."""

import argparse
import json
import math
from pathlib import Path
import sys

try:
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    from pptx.util import Pt
except ImportError as exc:
    raise SystemExit(
        "python-pptx 설치가 필요합니다: python -m pip install -r requirements.txt"
    ) from exc


def number(value, label, positive=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{label}: 숫자가 필요합니다.")
    if not math.isfinite(value) or (positive and value <= 0):
        raise ValueError(f"{label}: 올바른 유한 숫자가 필요합니다.")
    return value


def rectangle(record):
    pos = record["position_cm"]
    return tuple(number(pos[key], key, positive=key in ("width", "height"))
                 for key in ("x", "y", "width", "height"))


def cm(value):
    # Round instead of truncating float errors introduced by rotation inversion.
    return round(value * 360000)


def image_frame(record, warnings):
    """기존 추출기의 회전 후 외접 사각형에서 회전 전 프레임을 계산한다."""
    x, y, bw, bh = rectangle(record)
    angle = number(record.get("rotation_degrees", 0), "rotation_degrees") % 360
    c = abs(math.cos(math.radians(angle)))
    s = abs(math.sin(math.radians(angle)))
    determinant = c * c - s * s
    if abs(determinant) > 1e-6:
        width = (c * bw - s * bh) / determinant
        height = (c * bh - s * bw) / determinant
        if width > 0 and height > 0:
            return (x + (bw - width) / 2, y + (bh - height) / 2,
                    width, height, angle)
    warnings.append(
        f"도형 {record['shape_id']}: 회전 전 크기를 결정할 수 없어 "
        "외접 사각형에 회전 없이 배치했습니다."
    )
    return x, y, bw, bh, 0


def resource_path(folder, relative):
    path = (folder / relative).resolve()
    if not path.is_relative_to(folder.resolve()):
        raise ValueError(f"이미지 경로는 추출 폴더 내부여야 합니다: {relative}")
    if not path.is_file():
        raise FileNotFoundError(f"이미지를 찾을 수 없습니다: {path}")
    return path


def add_text(slide, record, font_name, body_size, title_size):
    x, y, width, height = rectangle(record)
    shape = slide.shapes.add_textbox(cm(x), cm(y), cm(width), cm(height))
    shape.name = f"text_{record['shape_id']}"
    if record["kind"] == "table_cell":
        shape.name += f"_r{record['row']}_c{record['column']}"
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.text = record["text"]
    for paragraph in frame.paragraphs:
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.font.name = font_name
        paragraph.font.size = Pt(title_size if record.get("is_title") else body_size)
        paragraph.font.bold = bool(record.get("is_title"))
    return shape


def add_image(slide, record, folder, warnings):
    path = resource_path(folder, record["file"])
    x, y, width, height, rotation = image_frame(record, warnings)
    shape = slide.shapes.add_picture(str(path), cm(x), cm(y), cm(width), cm(height))
    shape.name = f"image_{record['shape_id']}"
    shape.rotation = rotation
    for side, value in record.get("crop", {}).items():
        if side not in {"left", "top", "right", "bottom"}:
            raise ValueError(f"알 수 없는 이미지 자르기 방향: {side}")
        setattr(shape, f"crop_{side}", number(value, f"crop.{side}"))
    return shape


def rebuild_pptx(manifest, destination=None, font_name="맑은 고딕", body_size=18, title_size=28):
    """원본 PPTX를 열거나 복사하지 않는다. 입력은 추출 JSON과 이미지뿐이다."""
    manifest = Path(manifest)
    destination = Path(destination) if destination else manifest.parent / "copy.pptx"
    if destination.suffix.lower() != ".pptx":
        raise ValueError("출력 파일 확장자는 .pptx여야 합니다.")
    number(body_size, "body_size", positive=True)
    number(title_size, "title_size", positive=True)
    document = json.loads(manifest.read_text(encoding="utf-8-sig"))
    if document.get("coordinate_unit") != "cm" or document.get("coordinate_origin") != "top-left":
        raise ValueError("cm 단위, 왼쪽 위 원점의 content.json이 필요합니다.")
    prs = Presentation()
    size = document["slide_size_cm"]
    prs.slide_width = cm(number(size["width"], "slide width", positive=True))
    prs.slide_height = cm(number(size["height"], "slide height", positive=True))
    warnings = [
        "원본 글꼴·서식·배경·일반 도형·애니메이션은 추출 데이터에 없어 복원하지 않습니다.",
        "제외된 페이지 번호·머리말·꼬리말은 다시 생성하지 않습니다.",
        "텍스트는 외접 사각형 위치에 기본 서식으로 배치합니다. 텍스트 회전은 복원하지 않습니다.",
        "그룹 내부 이미지의 회전·뒤집기·기울임은 기존 추출 정보만으로 완전 복원할 수 없습니다.",
    ]
    report = {"input": manifest.name, "pages": 0, "texts": 0, "images": 0,
              "warnings": warnings}
    pages = sorted(document["slides"], key=lambda s: s["page"])
    if [page["page"] for page in pages] != list(range(1, len(pages) + 1)):
        raise ValueError("페이지 번호는 중복 없이 1부터 연속되어야 합니다.")
    table_cells = False
    fallback_order = False
    for page in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        records = [("image", r) for r in page["images"]]
        records += [("text", r) for r in page["texts"]]
        if all("z_order" in r for _, r in records):
            records.sort(key=lambda item: item[1]["z_order"])
        elif records:
            fallback_order = True
        for kind, record in records:
            if kind == "image":
                add_image(slide, record, manifest.parent, warnings)
                report["images"] += 1
            else:
                add_text(slide, record, font_name, body_size, title_size)
                table_cells |= record["kind"] == "table_cell"
                report["texts"] += 1
        report["pages"] += 1
    if table_cells:
        warnings.append("표 셀은 저장된 셀 위치의 텍스트 상자로 복원합니다. 표 테두리·빈 셀은 복원하지 않습니다.")
    if fallback_order:
        warnings.append("기존 JSON에 겹침 순서가 없어 이미지를 먼저, 텍스트를 위에 배치했습니다.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destination))
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("manifest", type=Path, nargs="?", default=Path("output/content.json"))
    parser.add_argument("-o", "--output", type=Path, help="출력 .pptx (기본: JSON 폴더/copy.pptx)")
    parser.add_argument("--font", default="맑은 고딕", help="대체 글꼴")
    parser.add_argument("--body-size", type=float, default=18, help="본문 크기(pt)")
    parser.add_argument("--title-size", type=float, default=28, help="제목 크기(pt)")
    args = parser.parse_args(argv)
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    destination = args.output if args.output else args.manifest.parent / "copy.pptx"
    try:
        report = rebuild_pptx(args.manifest, destination, args.font, args.body_size, args.title_size)
        destination.with_suffix(".report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except (OSError, ValueError, KeyError, TypeError, AttributeError) as exc:
        print(f"복원 실패: {exc}", file=sys.stderr)
        return 1
    print(f"생성 완료: {destination.resolve()} "
          f"(페이지 {report['pages']}, 텍스트 {report['texts']}, 이미지 {report['images']})")
    for warning in report["warnings"]:
        print(f"참고: {warning}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
